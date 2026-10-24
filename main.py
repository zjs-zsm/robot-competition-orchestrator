from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
import base64
import hashlib
import io
import json
import os
import random
import re
import sqlite3
import textwrap
import time
import urllib.request
import urllib.error
import urllib.parse
import zipfile

from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

APP_VERSION = "4.1.0"

app = FastAPI(
    title="Robot Competition Orchestrator",
    version=APP_VERSION,
    description="智能机器人创意竞赛助手 V4.1 动态创意/可修改报告中央编排器",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ============================================================
# 数据模型：保持与学习通插件现有字段兼容
# ============================================================
class ChatRequest(BaseModel):
    session_id: str = Field(..., description="会话ID")
    user_id: str = Field(..., description="用户ID")
    message: str = Field(..., description="用户输入")
    attachments: Optional[List[Dict[str, Any]]] = Field(default_factory=list)


class ChatResponse(BaseModel):
    success: bool
    stage: str
    intent: str
    message: str
    payload_json: str = ""
    download_url: str = ""
    data: Dict[str, Any] = Field(default_factory=dict)
    files: List[Dict[str, Any]] = Field(default_factory=list)
    suggested_actions: List[str] = Field(default_factory=list)
    version: int = 10
    updated_at: str


# ============================================================
# 基础配置
# ============================================================
PUBLIC_BASE_URL = os.getenv(
    "PUBLIC_BASE_URL",
    "https://robot-competition-orchestrator.onrender.com",
).rstrip("/")

DATA_DIR = Path(os.getenv("DATA_DIR", "/tmp/robot_competition_data"))
DATA_DIR.mkdir(parents=True, exist_ok=True)

EXPORT_DIR = Path(os.getenv("EXPORT_DIR", str(DATA_DIR / "exports")))
EXPORT_DIR.mkdir(parents=True, exist_ok=True)

SESSION_DB_PATH = os.getenv(
    "SESSION_DB_PATH",
    str(DATA_DIR / "sessions.sqlite3"),
)

# 可选：任意 OpenAI-compatible Chat Completions 服务。
# 没有配置时，系统会自动使用本地动态生成，不影响完整流程。
LLM_API_KEY = os.getenv("LLM_API_KEY", "").strip()
LLM_API_BASE = os.getenv("LLM_API_BASE", "").strip().rstrip("/")
LLM_MODEL = os.getenv("LLM_MODEL", "").strip()
LLM_TIMEOUT = int(os.getenv("LLM_TIMEOUT", "45"))

# 未来老师资料知识库接口：目前可留空。
KNOWLEDGE_API_URL = os.getenv("KNOWLEDGE_API_URL", "").strip()
KNOWLEDGE_API_KEY = os.getenv("KNOWLEDGE_API_KEY", "").strip()

# 可选图片生成/检索接口。留空时自动生成项目相关结构示意图。
IMAGE_API_URL = os.getenv("IMAGE_API_URL", "").strip()
IMAGE_API_KEY = os.getenv("IMAGE_API_KEY", "").strip()

# 推荐图像生成服务。配置后可自动生成“产品效果图 + 应用场景图”。
# 未配置时仍会生成本地概念插画与原生Word架构/流程可视化，不影响主流程。
POLLINATIONS_API_KEY = os.getenv("POLLINATIONS_API_KEY", "").strip()
POLLINATIONS_IMAGE_MODEL = os.getenv("POLLINATIONS_IMAGE_MODEL", "flux").strip() or "flux"
IMAGE_TIMEOUT = int(os.getenv("IMAGE_TIMEOUT", "90"))
STRICT_AI_IMAGES = os.getenv("STRICT_AI_IMAGES", "true").strip().lower() not in {"0", "false", "no", "off"}


# ============================================================
# 样例案例库：仅在没有真实知识库时用于流程验证
# ============================================================
REFERENCE_CASES = [
    {
        "title": "智能老人陪伴与健康监测机器人",
        "keywords": ["老人", "陪伴", "健康", "跌倒", "吃药", "家庭", "语音"],
        "track": "服务机器人",
    },
    {
        "title": "校园智能垃圾分类与回收机器人",
        "keywords": ["垃圾分类", "校园", "回收", "识别", "环保", "移动"],
        "track": "服务机器人",
    },
    {
        "title": "公共空间智能消毒巡检机器人",
        "keywords": ["消毒", "巡检", "公共空间", "环境", "导航", "安全"],
        "track": "特种机器人",
    },
    {
        "title": "导盲辅助与道路安全提醒机器人",
        "keywords": ["导盲", "道路", "安全", "避障", "语音", "辅助"],
        "track": "服务机器人",
    },
    {
        "title": "农业果蔬采摘与成熟度识别机器人",
        "keywords": ["农业", "采摘", "果蔬", "成熟度", "识别", "机械臂"],
        "track": "农业机器人",
    },
    {
        "title": "仓储物流自主搬运机器人",
        "keywords": ["物流", "仓储", "搬运", "路径规划", "导航", "调度"],
        "track": "工业机器人",
    },
    {
        "title": "水质检测与河道巡航机器人",
        "keywords": ["水质", "检测", "河道", "巡航", "传感器", "环保"],
        "track": "特种机器人",
    },
    {
        "title": "家庭教育陪伴与学习监督机器人",
        "keywords": ["儿童", "教育", "学习", "陪伴", "语音", "家庭"],
        "track": "服务机器人",
    },
]


# ============================================================
# 通用工具
# ============================================================
def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def normalize_text(text: str) -> str:
    return re.sub(r"\s+", "", (text or "").strip()).lower()


def clean_message(text: str) -> str:
    result = (text or "").strip()
    prefixes = [
        "启动V2中央编排器",
        "启动v2中央编排器",
        "启动机器人竞赛助手",
        "启动达标版竞赛助手",
        "运行机器人竞赛中央编排器",
    ]
    for prefix in prefixes:
        result = result.replace(prefix, "").strip()
    result = re.sub(r"/no_think\b", "", result, flags=re.I).strip()
    return result if result else (text or "").strip()


def contains_any(text: str, words: List[str]) -> bool:
    lowered = (text or "").lower()
    return any((word or "").lower() in lowered for word in words)


def clamp(value: float, low: float = 0, high: float = 100) -> float:
    return max(low, min(high, value))


def stable_seed(text: str) -> int:
    return int(hashlib.sha1(text.encode("utf-8")).hexdigest()[:8], 16)


def unique(items: List[str]) -> List[str]:
    return list(dict.fromkeys([x for x in items if x]))


def safe_filename(text: str, limit: int = 40) -> str:
    text = re.sub(r"[^\w\u4e00-\u9fff-]+", "_", text or "report")
    return text.strip("_")[:limit] or "report"


# ============================================================
# 会话持久化：SQLite，避免进程内变量导致短暂重启后完全丢失
# ============================================================
def _db() -> sqlite3.Connection:
    conn = sqlite3.connect(SESSION_DB_PATH, timeout=10)
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS sessions (
            session_key TEXT PRIMARY KEY,
            payload TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
        """
    )
    return conn


def build_session_key(req: ChatRequest) -> str:
    user_id = (req.user_id or "").strip()
    session_id = (req.session_id or "").strip()

    placeholder_users = {
        "", "anonymous", "current_user", "default_user", "user_456", "chaoxing-test-user"
    }
    placeholder_sessions = {
        "", "default", "default_session", "current_session", "current_session_123", "chaoxing-test-001"
    }

    if user_id in placeholder_users:
        user_id = "anonymous"
    if session_id in placeholder_sessions:
        session_id = "default"

    return f"{user_id}:{session_id}"


def load_session(session_key: str) -> Dict[str, Any]:
    try:
        with _db() as conn:
            row = conn.execute(
                "SELECT payload FROM sessions WHERE session_key = ?",
                (session_key,),
            ).fetchone()
        if row:
            return json.loads(row[0])
    except Exception as exc:
        print("SESSION_LOAD_ERROR:", repr(exc), flush=True)
    return {}


def save_session(session_key: str, session: Dict[str, Any]) -> None:
    try:
        payload = json.dumps(session, ensure_ascii=False)
        with _db() as conn:
            conn.execute(
                """
                INSERT INTO sessions(session_key, payload, updated_at)
                VALUES(?, ?, ?)
                ON CONFLICT(session_key) DO UPDATE SET
                    payload = excluded.payload,
                    updated_at = excluded.updated_at
                """,
                (session_key, payload, now_iso()),
            )
    except Exception as exc:
        print("SESSION_SAVE_ERROR:", repr(exc), flush=True)


# ============================================================
# 可选大模型调用
# ============================================================
def llm_available() -> bool:
    return bool(LLM_API_KEY and LLM_API_BASE and LLM_MODEL)


def call_llm_json(system_prompt: str, user_prompt: str) -> Optional[Dict[str, Any]]:
    if not llm_available():
        return None

    url = LLM_API_BASE
    if not url.endswith("/chat/completions"):
        url = url + "/chat/completions"

    body = {
        "model": LLM_MODEL,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        "temperature": 0.7,
        "response_format": {"type": "json_object"},
    }

    req = urllib.request.Request(
        url,
        data=json.dumps(body, ensure_ascii=False).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {LLM_API_KEY}",
            "Content-Type": "application/json",
        },
        method="POST",
    )

    try:
        with urllib.request.urlopen(req, timeout=LLM_TIMEOUT) as resp:
            raw = json.loads(resp.read().decode("utf-8"))
        content = raw["choices"][0]["message"]["content"]
        content = re.sub(r"^```(?:json)?|```$", "", content.strip(), flags=re.I).strip()
        return json.loads(content)
    except Exception as exc:
        print("LLM_CALL_ERROR:", repr(exc), flush=True)
        return None


# ============================================================
# 创意理解：本地动态解析 + 可选 LLM 增强
# ============================================================
SCENE_RULES = {
    "家庭/居家": ["家庭", "居家", "家里", "家用", "卧室", "客厅", "厨房", "阳台"],
    "校园/教育": ["校园", "学校", "教室", "学生", "宿舍", "实验室"],
    "医疗/养老": ["医院", "养老院", "老人", "康复", "护理", "患者"],
    "工业/仓储": ["工厂", "工业", "仓库", "物流", "产线", "搬运"],
    "农业/园艺": ["农业", "农田", "果园", "温室", "花盆", "植物", "多肉"],
    "环保/水域": ["水质", "河道", "湖泊", "环保", "垃圾", "污染"],
    "公共服务": ["社区", "商场", "车站", "公共", "道路", "景区"],
    "宠物照护": ["宠物", "猫", "狗", "猫砂", "喂食", "陪玩"],
}

TARGET_RULES = {
    "老年人": ["老人", "老年", "独居", "空巢"],
    "儿童/学生": ["儿童", "孩子", "学生", "小朋友"],
    "残障/特殊人群": ["残障", "盲人", "视障", "行动不便", "听障"],
    "宠物家庭": ["宠物", "猫", "狗"],
    "家庭用户": ["家庭", "家用", "居家", "家里"],
    "农业从业者": ["农民", "农业", "果园", "温室"],
    "工业/物流人员": ["工厂", "仓库", "物流", "工人"],
}

FUNCTION_RULES = {
    "用药提醒": ["吃药", "服药", "用药", "药物提醒"],
    "跌倒/姿态检测": ["跌倒", "摔倒", "姿态"],
    "语音交互与陪伴": ["语音", "聊天", "对话", "陪伴"],
    "健康监测": ["心率", "血压", "体温", "健康", "血氧"],
    "危险预警": ["危险", "报警", "异常", "烟雾", "燃气", "关火", "高温", "火灾"],
    "自主导航与避障": ["导航", "避障", "巡航", "移动", "跟随"],
    "视觉识别": ["识别", "视觉", "摄像头", "图像", "目标检测"],
    "机械执行": ["机械臂", "抓取", "搬运", "递送", "开关", "关闭燃气"],
    "远程通知": ["远程", "通知", "手机", "家属", "推送"],
    "智能家居联动": ["智能家居", "灯光", "空调", "门锁", "燃气阀"],
    "清洁整理": ["清洁", "扫地", "拖地", "收纳", "整理"],
    "宠物照护": ["猫砂", "补粮", "喂食", "陪玩", "宠物"],
    "植物养护": ["浇水", "植物", "多肉", "光照", "土壤", "施肥"],
    "垃圾分类与回收": ["垃圾分类", "回收", "垃圾"],
    "水质/环境检测": ["水质", "空气", "温湿度", "污染", "环境监测"],
    "学习辅助": ["学习", "作业", "教学", "辅导", "监督"],
}

TECH_RULES = {
    "多模态传感器融合": ["传感器", "跌倒", "健康", "烟雾", "高温", "燃气", "环境", "检测"],
    "计算机视觉识别": ["摄像头", "视觉", "图像", "识别", "目标检测", "姿态"],
    "语音识别与自然语言交互": ["语音", "聊天", "对话", "陪伴", "提醒"],
    "移动底盘与路径规划": ["移动", "导航", "避障", "巡航", "跟随"],
    "机械臂/执行机构控制": ["机械臂", "抓取", "搬运", "递送", "开关", "阀"],
    "物联网与远程通信": ["远程", "通知", "手机", "家属", "智能家居", "联网"],
    "边缘计算与本地决策": ["本地", "低延迟", "离线", "实时", "异常"],
    "智能规划与任务调度": ["自动", "规划", "调度", "定时", "任务"],
}


def _match_rules(text: str, rules: Dict[str, List[str]]) -> List[str]:
    return [name for name, keys in rules.items() if contains_any(text, keys)]


def _extract_free_keywords(raw: str) -> List[str]:
    cleaned = re.sub(r"[，。！？；、,.!?;:\n\t（）()\[\]【】]", " ", raw)
    parts = [p.strip() for p in cleaned.split() if 2 <= len(p.strip()) <= 12]
    stop = {"我想", "一个", "机器人", "智能机器人", "主要", "用于", "希望", "可以", "能够", "这个", "一种"}
    return unique([p for p in parts if p not in stop])[:12]


def local_extract_idea_fields(raw_idea: str) -> Dict[str, Any]:
    text = normalize_text(raw_idea)
    targets = _match_rules(text, TARGET_RULES) or ["目标场景用户"]
    scenes = _match_rules(text, SCENE_RULES) or ["实际应用场景"]
    functions = _match_rules(text, FUNCTION_RULES)
    tech = _match_rules(text, TECH_RULES)

    if not functions:
        functions = ["智能感知", "人机交互", "自主服务"]
    if not tech:
        tech = ["传感器采集与智能决策", "人机交互控制"]

    keywords = unique(
        _extract_free_keywords(raw_idea)
        + [k for name, keys in FUNCTION_RULES.items() for k in keys if k in text]
    )[:12]

    pain_points = []
    if contains_any(text, ["危险", "跌倒", "异常", "燃气", "火", "安全"]):
        pain_points.append("现有人工监护或被动告警存在发现不及时的问题")
    if contains_any(text, ["老人", "儿童", "残障", "宠物"]):
        pain_points.append("目标用户需要更低门槛、更主动的持续服务")
    if contains_any(text, ["自动", "提醒", "清洁", "搬运", "照护"]):
        pain_points.append("重复性任务占用时间，缺少持续自动化执行")
    if not pain_points:
        pain_points.append("现有方案智能化、连续服务和闭环执行能力不足")

    return {
        "raw_idea": raw_idea.strip(),
        "target_groups": unique(targets),
        "scenarios": unique(scenes),
        "core_functions": unique(functions),
        "tech_modules": unique(tech),
        "keywords": unique(keywords) or unique(functions)[:6],
        "pain_points": unique(pain_points),
        "design_goal": f"围绕{scenes[0]}中的真实需求，构建可感知、可判断、可执行、可反馈的机器人服务闭环。",
    }


def extract_idea_fields(raw_idea: str) -> Dict[str, Any]:
    local = local_extract_idea_fields(raw_idea)

    ai = call_llm_json(
        """你是大学生智能机器人创意竞赛方案分析专家。请只返回JSON，不要Markdown。
字段必须包含：target_groups(list), scenarios(list), core_functions(list), tech_modules(list),
keywords(list), pain_points(list), design_goal(str)。要求忠实于用户原始创意，不得强行改成老人陪护。""",
        f"用户创意：{raw_idea}\n请提取竞赛设计字段。",
    )

    if not ai:
        return local

    merged = dict(local)
    for key in ["target_groups", "scenarios", "core_functions", "tech_modules", "keywords", "pain_points"]:
        value = ai.get(key)
        if isinstance(value, list) and value:
            merged[key] = unique([str(x).strip() for x in value if str(x).strip()])[:10]
    if isinstance(ai.get("design_goal"), str) and ai["design_goal"].strip():
        merged["design_goal"] = ai["design_goal"].strip()
    merged["raw_idea"] = raw_idea.strip()
    return merged


# ============================================================
# 知识库接口：未来老师资料直接接这里
# ============================================================
def search_competition_knowledge(fields: Dict[str, Any], top_k: int = 5) -> Dict[str, Any]:
    query = " ".join(
        fields.get("keywords", [])
        + fields.get("core_functions", [])
        + fields.get("scenarios", [])
    )

    if KNOWLEDGE_API_URL:
        # Ping the lightweight health endpoint first: a free Render KB service may be asleep.
        # Never log the endpoint or API key; only log readiness fields.
        try:
            parsed_url = urllib.parse.urlsplit(KNOWLEDGE_API_URL)
            health_url = urllib.parse.urlunsplit((parsed_url.scheme, parsed_url.netloc, "/health", "", ""))
            health_req = urllib.request.Request(
                health_url,
                headers={"User-Agent": "RobotCompetitionAssistant/4.1"},
                method="GET",
            )
            with urllib.request.urlopen(health_req, timeout=8) as health_resp:
                health_data = json.loads(health_resp.read().decode("utf-8"))
            print(
                "KNOWLEDGE_API_HEALTH:",
                {"ok": health_data.get("ok"), "database_ready": health_data.get("database_ready")},
                flush=True,
            )
            if health_data.get("database_ready") is False:
                time.sleep(2)
        except Exception as exc:
            print("KNOWLEDGE_API_HEALTH_ERROR:", {"type": type(exc).__name__}, flush=True)
            time.sleep(2)

        body = json.dumps({"query": query, "top_k": top_k}, ensure_ascii=False).encode("utf-8")
        headers = {"Content-Type": "application/json"}
        if KNOWLEDGE_API_KEY:
            headers["Authorization"] = f"Bearer {KNOWLEDGE_API_KEY}"

        # Render free services may be waking from idle. Retry transient gateway,
        # timeout and connection errors rather than immediately falling back to samples.
        retryable_statuses = {408, 429, 500, 502, 503, 504}
        for attempt in range(1, 7):
            req = urllib.request.Request(
                KNOWLEDGE_API_URL,
                data=body,
                headers=headers,
                method="POST",
            )
            try:
                with urllib.request.urlopen(req, timeout=10) as resp:
                    data = json.loads(resp.read().decode("utf-8"))

                cases = data.get("cases") or data.get("results") or []
                if cases:
                    return {
                        "knowledge_base_used": True,
                        "similar_award_cases": cases[:top_k],
                        "knowledge_sources": data.get("sources", []),
                        "source_type": "external_knowledge_base",
                    }

                # A reachable API with no matching records is not evidence of real
                # award cases; preserve the existing honest sample-library fallback.
                print(
                    "KNOWLEDGE_API_EMPTY_RESULT:",
                    {"query_length": len(query), "attempt": attempt},
                    flush=True,
                )
                break

            except urllib.error.HTTPError as exc:
                print(
                    "KNOWLEDGE_API_HTTP_ERROR:",
                    {"status": exc.code, "attempt": attempt},
                    flush=True,
                )
                if exc.code not in retryable_statuses or attempt == 6:
                    break
            except (urllib.error.URLError, TimeoutError, ConnectionError) as exc:
                print(
                    "KNOWLEDGE_API_CONNECTION_ERROR:",
                    {"type": type(exc).__name__, "attempt": attempt},
                    flush=True,
                )
                if attempt == 6:
                    break
            except Exception as exc:
                print(
                    "KNOWLEDGE_API_ERROR:",
                    {"type": type(exc).__name__, "attempt": attempt},
                    flush=True,
                )
                break

            # Exponential backoff capped at 8s; allow enough time for a sleeping Render KB
            # service to complete its ~20-25s cold start before giving up.
            time.sleep(min(2 ** attempt, 8))

    # No usable records from the external API: sample cases are placeholders only.
    scored = []
    kws = set(fields.get("keywords", []))
    for case in REFERENCE_CASES:
        case_kws = set(case["keywords"])
        union = kws | case_kws
        sim = len(kws & case_kws) / len(union) if union else 0
        scored.append({**case, "similarity": round(sim, 3), "is_sample": True})
    scored.sort(key=lambda x: x["similarity"], reverse=True)
    return {
        "knowledge_base_used": False,
        "similar_award_cases": scored[:top_k],
        "knowledge_sources": [],
        "source_type": "sample_case_library",
    }

# ============================================================
# 动态候选题目
# ============================================================
def score_candidate(
    fields: Dict[str, Any],
    candidate: Dict[str, Any],
    knowledge: Dict[str, Any],
    candidate_index: int,
) -> Dict[str, Any]:
    function_count = len(fields.get("core_functions", []))
    tech_count = len(fields.get("tech_modules", []))
    keyword_count = len(fields.get("keywords", []))

    similar_cases = knowledge.get("similar_award_cases", [])
    highest_similarity = 0.0
    similar_case = {}
    if similar_cases:
        highest_similarity = float(similar_cases[0].get("similarity", 0) or 0)
        similar_case = similar_cases[0]

    # 不是实际获奖概率，是可解释竞争力评分
    innovation = 72 + min(12, function_count * 2.4) + candidate_index * 1.8
    scientific = 70 + min(15, tech_count * 2.8)
    application = 72 + min(14, len(fields.get("scenarios", [])) * 4) + min(5, len(fields.get("target_groups", [])) * 2)
    expression = 72 + min(12, keyword_count * 1.2)

    if "闭环" in candidate.get("positioning", ""):
        innovation += 3
    if candidate.get("differentiator"):
        innovation += 2

    similarity_penalty = highest_similarity * (12 if knowledge.get("knowledge_base_used") else 6)
    feasibility_penalty = 0
    if function_count >= 7 and tech_count <= 2:
        feasibility_penalty = 8
    elif function_count >= 6 and tech_count <= 2:
        feasibility_penalty = 5

    base = innovation * 0.35 + scientific * 0.30 + application * 0.25 + expression * 0.10
    competitiveness = clamp(base - similarity_penalty - feasibility_penalty)

    return {
        "innovation_score": round(clamp(innovation), 1),
        "scientific_score": round(clamp(scientific), 1),
        "application_score": round(clamp(application), 1),
        "expression_score": round(clamp(expression), 1),
        "base_total": round(base, 1),
        "highest_similarity": round(highest_similarity, 3),
        "similar_case": similar_case,
        "similarity_penalty": round(similarity_penalty, 1),
        "feasibility_penalty": round(feasibility_penalty, 1),
        "competitiveness_score": round(competitiveness, 1),
        "score_note": "可解释竞赛竞争力评分，不等同于真实获奖概率。",
    }


def _project_theme(fields: Dict[str, Any]) -> Tuple[str, str, str]:
    funcs = fields.get("core_functions", [])
    scene = fields.get("scenarios", ["应用场景"])[0]
    target = fields.get("target_groups", ["目标用户"])[0]

    if any("宠物" in f for f in funcs) or "宠物" in scene:
        return "宠护", "宠物照护", "智能照护"
    if any("植物" in f for f in funcs) or "园艺" in scene:
        return "植养", "植物养护", "智能园艺"
    if any("危险" in f for f in funcs) or any(x in fields.get("raw_idea", "") for x in ["燃气", "关火", "高温", "火灾"]):
        return "安防", "风险预警", "主动安全"
    if any("学习" in f for f in funcs):
        return "学伴", "学习辅助", "智能教育"
    if any("清洁" in f for f in funcs):
        return "净居", "家务服务", "自主清洁"
    if any("水质" in f or "环境" in f for f in funcs):
        return "环巡", "环境监测", "智能巡检"
    if any("机械" in f for f in funcs):
        return "智作", "任务执行", "自主作业"
    return "智服", f"{target}服务", f"{scene}智能服务"


def local_generate_candidates(fields: Dict[str, Any], knowledge: Dict[str, Any]) -> List[Dict[str, Any]]:
    seed = stable_seed(fields.get("raw_idea", ""))
    rng = random.Random(seed)
    prefix, domain, value = _project_theme(fields)
    target = fields.get("target_groups", ["目标用户"])[0]
    scene = fields.get("scenarios", ["实际应用场景"])[0]
    funcs = fields.get("core_functions", ["智能感知", "自主服务"])
    tech = fields.get("tech_modules", ["智能感知与决策"])
    function_text = "、".join(funcs[:3])
    tech_text = "、".join(tech[:3])

    name_stems = [
        ("星", "智"), ("卫", "慧"), ("伴", "灵"), ("盾", "安"),
        ("擎", "智"), ("航", "慧"), ("联", "云"), ("芯", "睿"),
    ]
    rng.shuffle(name_stems)

    names = [
        f"{prefix}{name_stems[0][0]}",
        f"{name_stems[1][1]}{domain[:2]}",
        f"{prefix}{name_stems[2][1]}",
    ]

    candidates = [
        {
            "id": 1,
            "title": f"“{names[0]}”——面向{target}的{domain}机器人",
            "positioning": f"围绕{function_text}形成感知—决策—执行—反馈闭环，突出实际需求解决能力。",
            "core_tech": tech_text,
            "differentiator": "场景需求闭环",
        },
        {
            "id": 2,
            "title": f"“{names[1]}”——基于多模态感知的{scene}{value}机器人",
            "positioning": f"以{tech_text}为核心，强调多源信息融合、主动判断与可靠执行。",
            "core_tech": tech_text,
            "differentiator": "多模态感知与主动决策",
        },
        {
            "id": 3,
            "title": f"“{names[2]}”——融合智能联动与自主服务的{domain}机器人",
            "positioning": f"突出{function_text}之间的协同联动，强调展示性、可扩展性和工程落地。",
            "core_tech": tech_text,
            "differentiator": "模块协同与可扩展设计",
        },
    ]

    for idx, item in enumerate(candidates):
        item["scores"] = score_candidate(fields, item, knowledge, idx)
    return candidates


def generate_candidates(fields: Dict[str, Any], knowledge: Dict[str, Any]) -> List[Dict[str, Any]]:
    knowledge_summary = [
        {"title": c.get("title"), "similarity": c.get("similarity", 0)}
        for c in knowledge.get("similar_award_cases", [])[:3]
    ]

    ai = call_llm_json(
        """你是中国大学生智能机器人创意竞赛命题专家。只返回JSON。
必须根据本次用户创意动态生成3个明显不同的题目，禁止固定输出“智护星/安居守护者/慧联家护”。
格式：{"candidates":[{"id":1,"title":"...","positioning":"...","core_tech":"...","differentiator":"..."},...]}。
题目要像正式竞赛作品名，技术路线可实现，三题侧重点不同。""",
        json.dumps({"fields": fields, "knowledge": knowledge_summary}, ensure_ascii=False),
    )

    if ai and isinstance(ai.get("candidates"), list) and len(ai["candidates"]) >= 3:
        result = []
        for idx, item in enumerate(ai["candidates"][:3]):
            candidate = {
                "id": idx + 1,
                "title": str(item.get("title", "")).strip(),
                "positioning": str(item.get("positioning", "")).strip(),
                "core_tech": str(item.get("core_tech", "")).strip() or "、".join(fields.get("tech_modules", [])[:3]),
                "differentiator": str(item.get("differentiator", "")).strip(),
            }
            if not candidate["title"]:
                return local_generate_candidates(fields, knowledge)
            candidate["scores"] = score_candidate(fields, candidate, knowledge, idx)
            result.append(candidate)
        return result

    return local_generate_candidates(fields, knowledge)


def format_candidate_message(fields: Dict[str, Any], candidates: List[Dict[str, Any]], prefix: str, knowledge: Dict[str, Any]) -> str:
    lines = [
        prefix,
        "",
        "【创意理解】",
        "目标用户：" + "、".join(fields.get("target_groups", [])),
        "应用场景：" + "、".join(fields.get("scenarios", [])),
        "核心功能：" + "、".join(fields.get("core_functions", [])),
        "技术模块：" + "、".join(fields.get("tech_modules", [])),
        "核心痛点：" + "；".join(fields.get("pain_points", [])),
        "",
        "【3个动态候选题目】",
    ]

    for c in candidates:
        s = c["scores"]
        lines += [
            "",
            f"{c['id']}. {c['title']}",
            f"定位：{c['positioning']}",
            f"差异化：{c.get('differentiator', '')}",
            f"核心技术：{c['core_tech']}",
            f"评分：创新性{s['innovation_score']}，科学性{s['scientific_score']}，应用前景{s['application_score']}，设计表达{s['expression_score']}",
            f"竞赛竞争力：{s['competitiveness_score']} / 100",
        ]

    if knowledge.get("knowledge_base_used"):
        lines += ["", "已接入真实知识库，本轮评分已参考知识库检索结果。"]
    else:
        lines += ["", "说明：当前尚未接入老师提供的真实往届资料，相似案例仅用于流程验证；评分不是实际获奖概率。"]

    lines += ["", "请输入 1、2、3 选择题目；输入“重新生成”换一批；也可以继续补充你的创意要求。"]
    return "\n".join(lines)


# ============================================================
# 7页动态报告
# ============================================================
def _local_report(fields: Dict[str, Any], selected: Dict[str, Any], knowledge: Dict[str, Any]) -> Dict[str, Any]:
    """生成内容较完整的7页报告底稿。

    这一层必须在没有外部LLM时也能独立工作；若配置LLM，build_report_json会再做一次增强。
    """
    title = selected["title"]
    target_list = fields.get("target_groups", []) or ["目标用户"]
    scene_list = fields.get("scenarios", []) or ["实际应用场景"]
    funcs = fields.get("core_functions", []) or ["智能感知", "任务决策", "自主服务"]
    techs = fields.get("tech_modules", []) or ["传感器采集", "智能决策", "执行控制"]
    pain = fields.get("pain_points", []) or ["现有方案连续服务能力不足"]
    target = "、".join(target_list)
    scenes = "、".join(scene_list)
    func_text = "、".join(funcs[:6])
    tech_text = "、".join(techs[:6])
    raw_idea = fields.get("raw_idea", "")

    score = selected.get("scores", {}).get("competitiveness_score")
    positioning = selected.get("positioning", "")
    differentiator = selected.get("differentiator", "")

    pages = {
        "page_1": {
            "page_no": 1,
            "module": "封面",
            "title": title,
            "content": {
                "subtitle": "智能机器人创意竞赛 · 项目设计报告",
                "project_positioning": positioning,
                "design_statement": (
                    f"本项目面向{target}，聚焦{scenes}，以{func_text}为主要功能，"
                    f"通过{tech_text}构建从感知、判断到执行和反馈的完整机器人服务闭环。"
                ),
                "keywords": fields.get("keywords", [])[:8],
                "competition_score": score,
            },
        },
        "page_2": {
            "page_no": 2,
            "module": "设计背景",
            "title": "设计背景与用户需求",
            "content": {
                "background_paragraphs": [
                    f"随着家庭服务、智能感知与物联网技术的发展，{scenes}正在从单点智能设备逐步向能够主动感知、持续服务和协同执行的机器人系统演进。"
                    f"对于{target}而言，真正有价值的产品并不是简单增加一个提醒或遥控功能，而是能够在无人持续操作的情况下完成连续任务，并在异常时及时反馈。",
                    f"本项目来源于用户原始创意：{raw_idea}。围绕这一创意，方案将需求拆分为“感知对象—状态判断—任务决策—执行动作—结果反馈”五个环节，避免只做功能堆叠，强调完整服务闭环与竞赛展示效果。",
                    f"现有同类产品常见问题包括：功能相互割裂、对真实场景适应性不足、异常处理依赖人工以及长期运行稳定性不高。因此本项目将{func_text}进行协同设计，并把可量化测试指标纳入方案。",
                ],
                "pain_points": pain,
                "user_needs": [
                    f"低门槛：面向{target}，交互方式应直观，减少复杂设置。",
                    f"持续性：在{scenes}中支持长期、重复任务自动执行。",
                    "主动性：不仅等待用户指令，还能够识别关键事件并主动触发服务。",
                    "可靠性：关键操作应有状态确认、失败重试和人工兜底。",
                    "可扩展性：预留传感器、执行机构、联网设备和知识库接口。",
                ],
                "design_goals": [
                    fields.get("design_goal", ""),
                    "构建可感知、可判断、可执行、可反馈的机器人服务闭环。",
                    "兼顾竞赛展示性、工程可实现性、用户体验和后续扩展空间。",
                    "通过明确的功能指标和场景测试，让创新点可以被验证而不是停留在概念描述。",
                ],
            },
        },
        "page_3": {
            "page_no": 3,
            "module": "产品整体结构",
            "title": "机器人整体结构设计",
            "content": {
                "overview": (
                    f"机器人总体采用模块化分层架构。系统围绕{scenes}中的任务需求，把硬件感知、边缘计算、任务决策、执行机构、"
                    f"人机交互与远程服务统一到同一控制链路中。核心模块包括{func_text}。"
                ),
                "system_layers": [
                    "感知层：通过摄像头、环境传感器、状态传感器或用户输入采集场景信息，并完成数据时间同步与基础过滤。",
                    "认知决策层：对感知结果进行融合，完成目标/状态识别、异常判断、任务优先级排序和执行策略生成。",
                    "执行层：根据任务类型驱动移动底盘、机械执行机构、语音提示、喂食/清洁/开关等功能模块完成动作。",
                    "交互与联网层：面向用户提供提示、确认、任务进度、异常消息和远程信息同步，并记录关键运行日志。",
                ],
                "core_modules": funcs,
                "mechanical_concept": (
                    f"本体结构围绕{scenes}进行模块化设计。主体可采用移动底盘或固定式底座，根据实际任务配置摄像头、传感器阵列、"
                    "主控计算单元、执行机构与交互终端。模块之间采用标准化供电与通信接口，便于竞赛阶段快速替换和迭代。"
                ),
                "system_block_diagram": {
                    "inputs": ["用户指令", "场景/对象状态", "传感器数据"],
                    "perception": techs[:2] or ["多源感知", "数据预处理"],
                    "decision": ["状态识别", "任务规划", "安全判断"],
                    "execution": funcs[:4] or ["执行机构", "服务动作"],
                    "feedback": ["执行确认", "异常提醒", "远程同步"],
                },
                "design_principles": [
                    "重心与运动机构优先保证稳定性和安全性。",
                    "传感器布置兼顾视野、遮挡、维护和线缆走向。",
                    "高频使用模块采用快拆或抽拉结构，方便清洁与维护。",
                    "外观与结构突出项目主题，使评委能够快速理解功能分区。",
                ],
            },
        },
        "page_4": {
            "page_no": 4,
            "module": "软硬件功能设计",
            "title": "硬件与软件功能设计",
            "content": {
                "hardware": [
                    {"name": "感知单元", "description": f"围绕{func_text}配置视觉、距离、环境或状态传感器；采样频率与安装位置按场景需求确定。"},
                    {"name": "主控计算单元", "description": "负责数据预处理、状态融合、任务调度与本地决策；关键基础功能在网络异常时仍可工作。"},
                    {"name": "执行机构", "description": "根据任务配置移动底盘、机械臂、舵机、电机、泵阀、喂食/清洁等专用机构，并增加限位和过流保护。"},
                    {"name": "通信与供电", "description": "通过Wi-Fi/Bluetooth/蜂窝网络或局域网与移动端、云端或智能家居设备交互，并配置电量检测与低电量策略。"},
                ],
                "software": [
                    {"name": "感知服务", "description": f"完成{tech_text}相关数据采集、预处理与质量检测。"},
                    {"name": "状态识别", "description": "将单一传感器结果转化为可用于决策的状态标签，并对不确定结果保留置信度。"},
                    {"name": "任务调度", "description": f"根据{func_text}设置优先级、触发条件、互斥条件和失败重试策略。"},
                    {"name": "用户交互", "description": "提供状态提示、任务确认、异常提醒、远程查看与历史记录，减少用户操作负担。"},
                ],
                "hardware_software_block": {
                    "input_side": ["视觉/环境/状态传感", "用户/移动端输入"],
                    "control_side": ["主控计算", "状态识别", "任务调度", "安全策略"],
                    "output_side": funcs[:4] or ["执行机构", "反馈终端"],
                },
                "interaction_logic": "用户/环境事件 → 多源感知 → 数据预处理 → 状态识别 → 风险与任务判断 → 执行动作 → 结果确认 → 记录与远程反馈",
                "functional_matrix": [
                    {"function": f, "trigger": "定时/事件/用户指令", "feedback": "状态提示 + 执行结果 + 异常记录"}
                    for f in funcs[:6]
                ],
            },
        },
        "page_5": {
            "page_no": 5,
            "module": "关键技术",
            "title": "关键技术与实现路线",
            "content": {
                "technical_summary": (
                    f"项目技术路线以{tech_text}为核心。实现时不追求复杂算法堆叠，而是把每项技术与具体任务、输入数据、判断逻辑、"
                    "执行动作和测试指标对应起来，保证方案能够从概念顺利过渡到样机。"
                ),
                "key_technologies": [
                    {
                        "name": t,
                        "implementation": (
                            f"围绕“{t}”设计输入数据、预处理方式、阈值/模型判断、输出动作和失败兜底。"
                            "竞赛阶段优先采用可复现、可解释、可快速迭代的实现方案，并记录关键日志用于调参。"
                        ),
                    }
                    for t in techs[:6]
                ],
                "technical_route": ["数据采集", "质量检查与预处理", "状态/目标识别", "多源信息融合", "任务决策", "执行控制", "反馈确认与日志记录"],
                "project_workflow": [
                    "场景触发/用户指令",
                    "多源感知采集",
                    "数据预处理",
                    "状态/目标识别",
                    "任务决策与安全判断",
                    "执行机构动作",
                    "结果确认与远程反馈",
                ],
                "engineering_metrics": [
                    "识别类：准确率、召回率、误报率与复杂光照/遮挡条件下稳定性。",
                    "实时类：感知到动作的端到端响应延迟。",
                    "执行类：任务执行成功率、重复定位误差、机构卡滞/失败率。",
                    "通信类：异常消息送达率、断网后的本地降级能力。",
                    "可靠性：连续运行时长、异常恢复能力和关键模块故障隔离。",
                ],
                "risk_control": [
                    "涉及安全的动作设置人工确认、软硬件限位和紧急停止。",
                    "网络异常时保留本地基础任务，恢复连接后再同步日志。",
                    "模型输出不确定时进入保守策略，避免直接执行高风险动作。",
                    "涉及图像、语音与用户信息时设置最小化采集、权限管理和本地化存储策略。",
                ],
            },
        },
        "page_6": {
            "page_no": 6,
            "module": "项目创新点",
            "title": "项目创新点与差异化分析",
            "content": {
                "innovation_intro": (
                    f"本项目的创新重点不是单独增加一个新功能，而是针对{target}在{scenes}中的连续需求，"
                    f"把{func_text}组成可协同、可验证、可扩展的机器人服务系统。"
                ),
                "innovation_points": [
                    {"name": "需求驱动", "description": f"从{target}在{scenes}中的真实痛点出发，功能设计与场景任务一一对应。"},
                    {"name": "闭环服务", "description": "把感知、判断、执行、确认和记录连接起来，避免停留在提醒或遥控层面。"},
                    {"name": "多模块协同", "description": f"让{'、'.join(funcs[:4])}按任务优先级协同工作，而不是相互独立展示。"},
                    {"name": "工程化设计", "description": "采用模块化软硬件架构、失败重试、状态确认与安全兜底，提高样机可实现性。"},
                    {"name": "持续演进", "description": "预留知识库、智能家居/移动端接口和新传感器接入能力，便于后续迭代。"},
                ],
                "comparison": [
                    {"dimension": "需求理解", "common": "以功能清单为主", "ours": "从场景任务和用户痛点反推功能"},
                    {"dimension": "交互方式", "common": "用户主动下达指令", "ours": "主动感知 + 事件触发 + 用户确认"},
                    {"dimension": "系统结构", "common": "模块相互独立", "ours": "统一任务调度与状态闭环"},
                    {"dimension": "可靠性", "common": "重演示、轻异常处理", "ours": "加入失败重试、日志和安全兜底"},
                    {"dimension": "扩展能力", "common": "功能固定", "ours": "模块接口化并预留知识库与外部设备接口"},
                ],
                "differentiation_analysis": {
                    "knowledge_base_used": knowledge.get("knowledge_base_used", False),
                    "similar_award_cases": knowledge.get("similar_award_cases", [])[:3],
                    "current_boundary": "未接入老师提供的真实获奖资料时，不把样例相似度宣传为真实获奖证据。",
                    "selected_differentiator": differentiator,
                },
            },
        },
        "page_7": {
            "page_no": 7,
            "module": "行业应用前景",
            "title": "行业应用前景与落地路径",
            "content": {
                "prospect_paragraphs": [
                    f"从应用价值看，项目面向{target}，可首先在{scenes}完成小规模验证。与单一智能设备相比，机器人具备移动/执行、持续感知和多任务协同能力，"
                    "因此更适合承担具有连续性、重复性和异常处置要求的服务任务。",
                    "从产品化路径看，应优先完成核心任务闭环，再逐步增加高级功能。竞赛样机阶段重点证明功能可行和场景价值；后续可通过模块化版本、移动端服务、"
                    "智能家居或行业平台接入形成更完整的产品体系。",
                ],
                "application_scenarios": scene_list,
                "target_users": target_list,
                "landing_roadmap_title": "从竞赛样机到场景落地",
                "deployment_paths": [
                    "阶段1：完成核心功能最小可行样机，验证主任务闭环。",
                    "阶段2：在目标场景开展连续测试，记录成功、失败和异常样本。",
                    "阶段3：根据测试结果优化结构、算法、交互和安全策略。",
                    "阶段4：形成模块化版本，适配不同家庭/场景和成本档位。",
                    "阶段5：接入移动端、智能家居、知识库或行业平台，构建持续服务能力。",
                ],
                "social_value": [
                    "降低重复性人工工作负担，提高用户时间利用效率。",
                    "提高异常发现和信息反馈速度，减少关键事件漏报。",
                    "通过持续服务与低门槛交互提升目标用户使用体验。",
                    "形成可扩展的数据与服务闭环，为后续算法迭代提供基础。",
                ],
                "future_iterations": [
                    "接入老师提供的真实往届获奖作品知识库，完善相似度与竞争力分析。",
                    "增加真实场景测试数据和可量化指标，形成更有说服力的竞赛证据。",
                    "优化机器人外观、结构与交互细节，提升作品完成度和展示效果。",
                    "补齐答辩PPT、演示视频、结构图纸和测试记录，形成完整参赛材料。",
                ],
            },
        },
    }

    return {
        "schema_version": "2.0",
        "project_title": title,
        "raw_idea": raw_idea,
        "selected_candidate": selected,
        "idea_fields": fields,
        "knowledge": {
            "knowledge_base_used": knowledge.get("knowledge_base_used", False),
            "similar_award_cases": knowledge.get("similar_award_cases", [])[:5],
            "similarity_score": (knowledge.get("similar_award_cases") or [{}])[0].get("similarity") if knowledge.get("similar_award_cases") else None,
            "competition_probability": None,
            "knowledge_sources": knowledge.get("knowledge_sources", []),
        },
        "pages": pages,
        "generated_at": now_iso(),
    }

def build_report_json(fields: Dict[str, Any], selected: Dict[str, Any], knowledge: Dict[str, Any]) -> Dict[str, Any]:
    base = _local_report(fields, selected, knowledge)

    ai = call_llm_json(
        """你是智能机器人竞赛设计报告专家。只返回JSON。
必须保持严格7页：page_1封面、page_2设计背景、page_3整体结构、page_4软硬件功能、page_5关键技术、page_6创新点、page_7应用前景。
内容必须完全围绕给定用户创意和已选题目，禁止套用老人陪护固定模板。每页内容要专业、具体、可落地。
返回格式：{"pages":{...}}，不要删除页码/module/title/content字段。""",
        json.dumps({"base_report": base, "fields": fields, "selected": selected}, ensure_ascii=False),
    )

    if ai and isinstance(ai.get("pages"), dict):
        # 只替换7页内容，保留元数据和知识库字段
        for key in [f"page_{i}" for i in range(1, 8)]:
            if isinstance(ai["pages"].get(key), dict):
                page = ai["pages"][key]
                page.setdefault("page_no", int(key.split("_")[1]))
                page.setdefault("module", base["pages"][key]["module"])
                page.setdefault("title", base["pages"][key]["title"])
                page.setdefault("content", base["pages"][key]["content"])
                base["pages"][key] = page
    return base


def format_report_summary(report_json: Dict[str, Any]) -> str:
    lines = [
        "7页结构化竞赛报告已经生成完成。",
        "",
        f"项目题目：{report_json.get('project_title', '')}",
        "",
        "报告固定结构：",
    ]
    for i in range(1, 8):
        page = report_json["pages"][f"page_{i}"]
        lines.append(f"{i}. {page.get('title', page.get('module', ''))}")
    lines += [
        "",
        "结构化数据已保存在 payload_json 中。",
        "输入“下载报告”即可生成图文并茂的 Word 文档。",
    ]
    if not report_json.get("knowledge", {}).get("knowledge_base_used"):
        lines += ["", "当前尚未接入老师提供的真实获奖作品知识库；接口已预留，后续直接接入即可。"]
    return "\n".join(lines)


# ============================================================
# 动态图片：项目相关概念图、结构图、流程图、创新图
# ============================================================
def _domain_hint(report: Dict[str, Any]) -> str:
    fields = report.get("idea_fields", {})
    raw = fields.get("raw_idea", "")
    text = normalize_text(raw + " " + " ".join(fields.get("core_functions", [])))
    if contains_any(text, ["宠物", "猫", "狗", "猫砂", "喂食"]):
        return "pet care robot in a modern home, cat or dog nearby"
    if contains_any(text, ["老人", "养老", "跌倒", "用药"]):
        return "home elder-care companion robot assisting an older adult"
    if contains_any(text, ["厨房", "燃气", "关火", "烹饪"]):
        return "home kitchen safety robot near a stove and cooking area"
    if contains_any(text, ["植物", "浇水", "多肉", "园艺"]):
        return "home plant-care robot tending potted plants on a balcony"
    if contains_any(text, ["儿童", "学习", "教育", "作业"]):
        return "family educational companion robot helping a child study"
    if contains_any(text, ["清洁", "扫地", "拖地", "收纳"]):
        return "home service robot doing cleaning and organizing"
    if contains_any(text, ["农业", "果园", "温室", "采摘"]):
        return "agricultural service robot working in a greenhouse or orchard"
    if contains_any(text, ["仓库", "物流", "搬运"]):
        return "autonomous warehouse service robot moving goods"
    return "smart service robot operating in its intended real-world environment"


def _visual_prompt(report: Dict[str, Any], kind: str) -> str:
    fields = report.get("idea_fields", {})
    title = report.get("project_title", "smart service robot")
    funcs = ", ".join(fields.get("core_functions", [])[:6])
    techs = ", ".join(fields.get("tech_modules", [])[:5])
    domain = _domain_hint(report)

    if kind == "product":
        return (
            f"PRODUCT DESIGN HERO IMAGE for university robotics competition project: {title}. "
            f"Design one realistic, BUILDABLE engineering prototype robot for {domain}. "
            f"Functions: {funcs}. Technologies: {techs}. "
            "Show the complete robot body as the dominant subject, three-quarter front view, "
            "industrial design prototype, visible sensors, actuators, access panels and mobility mechanism. "
            "Neutral premium studio or exhibition pedestal background, controlled soft lighting, "
            "no pet beside the robot, no human, no room-wide action scene, no phone UI, "
            "NO text, NO labels, NO logo, NO watermark. Photorealistic product visualization, 16:9."
        )

    if kind == "scene":
        return (
            f"REAL APPLICATION SCENE for robotics competition project: {title}. "
            f"The robot must be visibly performing one concrete task in {domain}. Functions: {funcs}. "
            "Use a wide documentary composition in a real environment, side or diagonal camera angle, "
            "show the target user/object and the task result clearly. The robot should occupy less than half of the frame. "
            "For pet-care projects, show the robot interacting with a cat/dog, feeder/litter area or monitoring activity; "
            "for elder-care, show safe assistance; for kitchen safety, show detection/control near the stove; "
            "for plant-care, show watering/inspection beside plants. "
            "This image MUST look obviously different from a studio product render: different composition, different background, "
            "active task, environmental context, realistic photography. NO text, NO logo, NO watermark, 16:9."
        )

    return (
        f"Functional engineering detail of robot project {title}; sensors and actuators for {funcs}; "
        f"technologies {techs}; realistic prototype, no text, no logo, 16:9."
    )

def _download_binary(url: str, dest: Path, headers: Optional[Dict[str, str]] = None, timeout: int = 90) -> bool:
    try:
        req = urllib.request.Request(url, headers=headers or {}, method="GET")
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            data = resp.read()
        if len(data) < 5000:
            return False
        dest.write_bytes(data)
        return True
    except Exception as exc:
        print("IMAGE_DOWNLOAD_ERROR:", repr(exc), flush=True)
        return False


def _pollinations_image(prompt: str, dest: Path) -> bool:
    """Pollinations 图片生成。优先 POST OpenAI-compatible 接口，失败后 GET 兜底。"""
    if not POLLINATIONS_API_KEY:
        print("POLLINATIONS_ERROR: API key missing", flush=True)
        return False

    # 方式1：OpenAI-compatible 图片生成接口
    try:
        url = "https://gen.pollinations.ai/v1/images/generations"
        payload = {
            "model": POLLINATIONS_IMAGE_MODEL or "flux",
            "prompt": prompt,
            "size": "1536x1024",
            "n": 1,
            "response_format": "b64_json",
        }
        req = urllib.request.Request(
            url,
            data=json.dumps(payload).encode("utf-8"),
            headers={
                "Authorization": f"Bearer {POLLINATIONS_API_KEY}",
                "Content-Type": "application/json",
                "Accept": "application/json",
            },
            method="POST",
        )
        with urllib.request.urlopen(req, timeout=IMAGE_TIMEOUT) as resp:
            raw = resp.read()
            status = getattr(resp, "status", 200)
        print(f"POLLINATIONS_POST_STATUS: {status}", flush=True)
        data = json.loads(raw.decode("utf-8"))
        items = data.get("data") or []
        if items:
            item = items[0]
            if item.get("b64_json"):
                image_bytes = base64.b64decode(item["b64_json"])
                dest.write_bytes(image_bytes)
                if dest.exists() and dest.stat().st_size > 5000:
                    print(f"POLLINATIONS_IMAGE_SUCCESS: {dest} {dest.stat().st_size} bytes", flush=True)
                    return True
            if item.get("url") and _download_binary(item["url"], dest, timeout=IMAGE_TIMEOUT):
                print(f"POLLINATIONS_IMAGE_SUCCESS_URL: {dest}", flush=True)
                return True
    except urllib.error.HTTPError as exc:
        try:
            body = exc.read().decode("utf-8", errors="ignore")
        except Exception:
            body = ""
        print(f"POLLINATIONS_POST_HTTP_ERROR: {exc.code} {exc.reason} {body[:1000]}", flush=True)
    except Exception as exc:
        print("POLLINATIONS_POST_ERROR:", repr(exc), flush=True)

    # 方式2：GET /image/{prompt} 兜底；key 放 query，避免部分代理丢 Authorization
    try:
        encoded = urllib.parse.quote(prompt[:1800], safe="")
        query = urllib.parse.urlencode({
            "model": POLLINATIONS_IMAGE_MODEL or "flux",
            "width": 1280,
            "height": 768,
            "nologo": "true",
            "key": POLLINATIONS_API_KEY,
        })
        url = f"https://gen.pollinations.ai/image/{encoded}?{query}"
        req = urllib.request.Request(
            url,
            headers={"User-Agent": "RobotCompetitionAssistant/4.0", "Accept": "image/*"},
        )
        with urllib.request.urlopen(req, timeout=IMAGE_TIMEOUT) as resp:
            image_bytes = resp.read()
            content_type = (resp.headers.get("Content-Type") or "").lower()
        print(f"POLLINATIONS_GET_STATUS: {content_type}, {len(image_bytes)} bytes", flush=True)
        if len(image_bytes) > 5000 and (
            "image/" in content_type
            or image_bytes[:3] == b"\xff\xd8\xff"
            or image_bytes[:8] == b"\x89PNG\r\n\x1a\n"
        ):
            dest.write_bytes(image_bytes)
            print(f"POLLINATIONS_GET_SUCCESS: {dest}", flush=True)
            return True
    except urllib.error.HTTPError as exc:
        try:
            body = exc.read().decode("utf-8", errors="ignore")
        except Exception:
            body = ""
        print(f"POLLINATIONS_GET_HTTP_ERROR: {exc.code} {exc.reason} {body[:1000]}", flush=True)
    except Exception as exc:
        print("POLLINATIONS_GET_ERROR:", repr(exc), flush=True)
    return False

def _commons_search_terms(report: Dict[str, Any], kind: str) -> str:
    hint = _domain_hint(report)
    if "pet care" in hint:
        return "pet care robot cat feeder robot" if kind == "product" else "cat home smart pet feeder"
    if "elder-care" in hint:
        return "service robot elder care" if kind == "product" else "elderly home service robot"
    if "kitchen" in hint:
        return "kitchen service robot" if kind == "product" else "smart kitchen robot"
    if "plant-care" in hint:
        return "agricultural robot plant care" if kind == "product" else "smart irrigation potted plants"
    if "educational" in hint:
        return "educational robot" if kind == "product" else "robot child education"
    if "cleaning" in hint:
        return "home cleaning robot" if kind == "product" else "robot vacuum home"
    if "agricultural" in hint:
        return "agricultural robot" if kind == "product" else "farm robot greenhouse"
    if "warehouse" in hint:
        return "warehouse robot" if kind == "product" else "autonomous mobile robot warehouse"
    return "service robot prototype" if kind == "product" else "service robot home"


def _wikimedia_commons_image(report: Dict[str, Any], kind: str, dest: Path) -> bool:
    """无需API Key的真实图片兜底。仅从Wikimedia Commons开放资源检索。"""
    query = _commons_search_terms(report, kind)
    params = {
        "action": "query",
        "generator": "search",
        "gsrsearch": query,
        "gsrnamespace": "6",
        "gsrlimit": "8",
        "prop": "imageinfo",
        "iiprop": "url|mime",
        "format": "json",
        "origin": "*",
    }
    url = "https://commons.wikimedia.org/w/api.php?" + urllib.parse.urlencode(params)
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "RobotCompetitionAssistant/2.0"})
        with urllib.request.urlopen(req, timeout=20) as resp:
            data = json.loads(resp.read().decode("utf-8"))
        pages = list((data.get("query", {}).get("pages", {}) or {}).values())
        for page in pages:
            infos = page.get("imageinfo") or []
            if not infos:
                continue
            info = infos[0]
            mime = str(info.get("mime", "")).lower()
            image_url = info.get("url")
            if image_url and mime in {"image/jpeg", "image/png", "image/webp"}:
                if _download_binary(image_url, dest, headers={"User-Agent": "RobotCompetitionAssistant/2.0"}, timeout=30):
                    return True
    except Exception as exc:
        print("WIKIMEDIA_IMAGE_ERROR:", repr(exc), flush=True)
    return False


def _generic_image_api(prompt: str, dest: Path) -> bool:
    """兼容用户未来自建/其他图片服务。
    约定POST JSON {prompt,width,height}，返回图片二进制，或JSON中的url字段。
    """
    if not IMAGE_API_URL:
        return False
    body = json.dumps({"prompt": prompt, "width": 1280, "height": 768}, ensure_ascii=False).encode("utf-8")
    headers = {"Content-Type": "application/json"}
    if IMAGE_API_KEY:
        headers["Authorization"] = f"Bearer {IMAGE_API_KEY}"
    try:
        req = urllib.request.Request(IMAGE_API_URL, data=body, headers=headers, method="POST")
        with urllib.request.urlopen(req, timeout=IMAGE_TIMEOUT) as resp:
            ctype = (resp.headers.get("Content-Type") or "").lower()
            data = resp.read()
        if "image" in ctype and len(data) > 5000:
            dest.write_bytes(data)
            return True
        raw = json.loads(data.decode("utf-8"))
        image_url = raw.get("url") or raw.get("image_url") or raw.get("data", {}).get("url")
        if image_url:
            return _download_binary(image_url, dest, timeout=IMAGE_TIMEOUT)
    except Exception as exc:
        print("GENERIC_IMAGE_API_ERROR:", repr(exc), flush=True)
    return False


def _fallback_robot_illustration(report: Dict[str, Any], dest: Path, scene: bool = False) -> None:
    """高质量离线概念插画。无外部图片服务时也生成可嵌入Word的产品图/校园场景图。"""
    W, H = 1280, 768
    img = Image.new("RGB", (W, H), (242, 246, 252))
    d = ImageDraw.Draw(img)
    font = ImageFont.load_default()

    def gradient(top, bottom):
        for y in range(H):
            t = y / max(1, H - 1)
            color = tuple(int(top[i] * (1 - t) + bottom[i] * t) for i in range(3))
            d.line((0, y, W, y), fill=color)

    def rounded(box, radius, fill, outline=None, width=1):
        d.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)

    def robot(cx, cy, scale=1.0, night=False):
        def box(x1, y1, x2, y2):
            return (int(cx + x1 * scale), int(cy + y1 * scale),
                    int(cx + x2 * scale), int(cy + y2 * scale))
        outline = (45, 66, 91) if not night else (29, 44, 62)
        # soft ground shadow and rugged wheels
        d.ellipse(box(-205, 115, 205, 205), fill=(32, 45, 62) if night else (207, 216, 228))
        for x in (-145, 118):
            d.ellipse(box(x-38, 65, x+38, 170), fill=(31, 39, 50), outline=(13, 19, 28), width=max(2,int(4*scale)))
            d.ellipse(box(x-20, 83, x+20, 151), fill=(111, 127, 145), outline=(220, 228, 238), width=max(1,int(3*scale)))
            d.ellipse(box(x-8, 96, x+8, 138), fill=(45, 57, 73))
        # lower chassis, front bumper and shell
        d.polygon([box(-190, 20, -130, 105)[0:2], box(-130, 105, 170, 140)[2:4],
                   box(190, 15, 150, 100)[0:2], box(-160, -5, -190, 30)[0:2]], fill=(62, 83, 109))
        rounded(box(-190, -15, 190, 112), int(34*scale), (211, 222, 237), outline, max(2,int(5*scale)))
        rounded(box(-154, 0, 155, 78), int(24*scale), (236, 242, 249), (129, 151, 177), max(1,int(3*scale)))
        # front sensor panel
        rounded(box(-90, 18, 95, 73), int(14*scale), (48, 66, 89), (21, 34, 49), max(1,int(3*scale)))
        d.ellipse(box(-63, 29, -31, 61), fill=(30, 38, 51), outline=(92, 173, 232), width=max(2,int(4*scale)))
        d.ellipse(box(-54, 36, -40, 50), fill=(91, 196, 245))
        d.ellipse(box(24, 31, 43, 50), fill=(74, 198, 255))
        d.ellipse(box(51, 31, 70, 50), fill=(255, 91, 85))
        # upper compute enclosure / mast
        rounded(box(-74, -104, 76, -17), int(20*scale), (224, 232, 243), outline, max(2,int(4*scale)))
        rounded(box(-56, -91, 58, -38), int(12*scale), (52, 72, 97), (105, 132, 162), max(1,int(3*scale)))
        d.ellipse(box(-35, -82, 2, -45), fill=(21, 30, 43), outline=(99, 202, 255), width=max(2,int(4*scale)))
        d.ellipse(box(-25, -73, -7, -55), fill=(74, 182, 241))
        # 360-degree lidar and antennas
        rounded(box(-94, -132, 94, -105), int(10*scale), (79, 103, 132), outline, max(1,int(3*scale)))
        d.line(box(-80, -119, 80, -119), fill=(99, 221, 255), width=max(2,int(4*scale)))
        d.ellipse(box(-12, -144, 12, -120), fill=(40, 53, 70), outline=(122, 224, 255), width=max(1,int(3*scale)))
        d.line(box(60, -102, 88, -156), fill=(75, 103, 133), width=max(2,int(4*scale)))
        d.ellipse(box(80, -165, 96, -149), fill=(80, 190, 255))
        # proximity waves / headlamp
        ray = (110, 205, 255) if night else (87, 157, 219)
        for k in range(3):
            pad = (k + 1) * 17
            d.arc(box(-125-pad, -180-pad, 125+pad, 10+pad), 195, 345, fill=ray, width=max(1,int(3*scale)))
        rounded(box(-174, 84, 174, 105), int(8*scale), (76, 98, 124), None, 1)
        return

    if scene:
        gradient((13, 24, 48), (38, 68, 91))
        # moon and stars
        d.ellipse((1030, 60, 1110, 140), fill=(245, 235, 194))
        d.ellipse((1055, 48, 1125, 118), fill=(20, 39, 65))
        for x, y, r in [(120,90,3),(240,150,2),(430,80,3),(650,125,2),(820,70,3),(950,190,2),(1150,205,2),(535,180,2)]:
            d.ellipse((x-r,y-r,x+r,y+r), fill=(222, 237, 255))
        # campus buildings with lit windows
        for x, y, w, h in [(30,240,240,300),(290,185,230,355),(850,230,190,310),(1060,275,180,265)]:
            d.rectangle((x,y,x+w,y+h), fill=(31,48,72), outline=(82,111,145), width=3)
            d.rectangle((x+12,y+12,x+w-12,y+45), fill=(47,70,97))
            for wx in range(x+22,x+w-16,42):
                for wy in range(y+65,y+h-20,55):
                    d.rectangle((wx,wy,wx+17,wy+25), fill=(246,207,120) if (wx+wy)//7%3 else (92,143,183))
        # trees and path
        for x, y in [(80,510),(245,520),(545,520),(790,520),(1120,515)]:
            d.rectangle((x-8,y-75,x+8,y+15), fill=(83,69,59))
            for dx,dy,rx in [(-18,-76,35),(12,-95,39),(0,-125,34)]:
                d.ellipse((x+dx-rx,y+dy-rx,x+dx+rx,y+dy+rx), fill=(33,94,83), outline=(57,131,106), width=2)
        d.polygon([(0,580),(1280,555),(1280,768),(0,768)], fill=(66,83,99))
        d.polygon([(480,768),(585,558),(740,558),(900,768)], fill=(103,119,132))
        d.line((660,768,665,580), fill=(236,206,125), width=5)
        for yy in range(610,760,45):
            d.line((661,yy,663,yy+20), fill=(242,224,163), width=4)
        # patrol robot in the foreground
        robot(660, 600, 0.70, night=True)
        # route markers and remote alert panel
        for x,y in [(390,680),(430,650),(470,620),(510,600)]:
            d.ellipse((x-7,y-7,x+7,y+7), fill=(84,220,255), outline=(220,248,255), width=2)
        rounded((1030,370,1210,650), 28, (237,244,251), (117,153,190), 5)
        rounded((1052,398,1188,432), 10, (48,68,94))
        d.ellipse((1083,455,1155,527), fill=(224,78,70))
        d.polygon([(1118,468),(1130,496),(1106,496)], fill=(255,255,255))
        d.rectangle((1083,545,1156,555), fill=(110,139,170))
        d.rectangle((1070,574,1170,584), fill=(177,193,210))
        d.rectangle((1090,604,1150,614), fill=(91,171,219))
        d.text((1058, 637), "ALERT SENT", fill=(45,70,98), font=font)
        d.text((55, 36), "CAMPUS NIGHT PATROL", fill=(226,239,251), font=font)
        d.text((55, 58), "EDGE AI  /  AUTONOMOUS NAVIGATION  /  REMOTE ALERT", fill=(156,194,224), font=font)
    else:
        gradient((239, 245, 253), (218, 229, 242))
        # studio halo and floor
        d.ellipse((300,70,980,700), fill=(232,240,250))
        d.ellipse((220,585,1050,730), fill=(200,213,230))
        d.ellipse((360,610,930,705), fill=(185,201,220))
        # technical accents
        d.rounded_rectangle((65,80,365,160), radius=18, fill=(250,252,255), outline=(174,193,214), width=2)
        d.text((88,103), "EDGE AI  /  STM32 CONTROL", fill=(48,72,101), font=font)
        d.rounded_rectangle((910,80,1210,160), radius=18, fill=(250,252,255), outline=(174,193,214), width=2)
        d.text((934,103), "CAMERA  /  LiDAR  /  SONAR", fill=(48,72,101), font=font)
        # soft sensor rays behind the rover
        for radius in (190,245,300):
            d.arc((640-radius,360-radius,640+radius,360+radius), 210, 330, fill=(128,173,213), width=3)
        robot(650, 475, 1.10, night=False)
        # floating sensor modules and callouts
        for x,y,fill in [(250,300,(75,153,213)),(1030,300,(61,176,183)),(270,500,(236,170,71)),(1010,500,(104,132,191))]:
            d.ellipse((x-30,y-30,x+30,y+30), fill=fill, outline=(255,255,255), width=4)
            d.ellipse((x-10,y-10,x+10,y+10), fill=(244,249,255))
        d.line((280,320,420,365), fill=(78,133,188), width=4)
        d.line((1000,320,850,360), fill=(61,150,164), width=4)
        d.line((300,500,455,490), fill=(204,146,59), width=4)
        d.line((980,500,850,495), fill=(91,115,167), width=4)
        d.text((60,690), "ROBOT CONCEPT  /  MODULAR SENSING  /  SAFE MOTION", fill=(54,78,106), font=font)

    img.save(dest, format="PNG", optimize=True)

def _image_signature(path: Path) -> Optional[List[int]]:
    """64-bit average hash，用来避免第1、2页生成几乎同一张图。"""
    try:
        with Image.open(path) as img:
            img = img.convert("L").resize((8, 8))
            vals = list(img.getdata())
        avg = sum(vals) / len(vals)
        return [1 if v >= avg else 0 for v in vals]
    except Exception:
        return None


def _image_similarity(path_a: Path, path_b: Path) -> float:
    a = _image_signature(path_a)
    b = _image_signature(path_b)
    if not a or not b or len(a) != len(b):
        return 0.0
    same = sum(1 for x, y in zip(a, b) if x == y)
    return same / len(a)


def _generate_ai_visual(prompt: str, path: Path) -> bool:
    """优先用户自建图片接口，其次 Pollinations。AI 已配置时不再静默塞入简笔画。"""
    return _generic_image_api(prompt, path) or _pollinations_image(prompt, path)


def generate_project_images(report: Dict[str, Any], session_key: str, fast_mode: bool = False) -> Dict[str, str]:
    digest = hashlib.sha1(f"{session_key}-{report.get('project_title','')}".encode("utf-8")).hexdigest()[:10]
    result: Dict[str, str] = {}
    ai_configured = bool(POLLINATIONS_API_KEY or IMAGE_API_URL)

    # 第1页：产品英雄图
    product_path = EXPORT_DIR / f"visual_product_{digest}.png"
    if fast_mode:
        # Chaoxing's workflow can time out while waiting for external image APIs.
        # Fast export uses local illustrative images; the Word-native system/workflow diagrams remain editable.
        _fallback_robot_illustration(report, product_path, scene=False)
        scene_path = EXPORT_DIR / f"visual_scene_{digest}.png"
        _fallback_robot_illustration(report, scene_path, scene=True)
        result["product"] = str(product_path)
        result["scene"] = str(scene_path)
        return result

    product_prompt = _visual_prompt(report, "product")
    product_ok = _generate_ai_visual(product_prompt, product_path) if ai_configured else False

    if not product_ok:
        if ai_configured and STRICT_AI_IMAGES:
            raise RuntimeError(
                "第1页AI产品效果图生成失败。为避免再次输出简笔画，已停止Word导出；请查看Render日志中的POLLINATIONS_*错误后重试。"
            )
        # 本地/无密钥测试环境仍允许开放图片兜底，便于验证Word主流程
        product_ok = _wikimedia_commons_image(report, "product", product_path)
        if not product_ok:
            _fallback_robot_illustration(report, product_path, scene=False)
    result["product"] = str(product_path)

    # 第2页：真实应用场景图。最多3次，若与封面过于相似则重生成。
    scene_path = EXPORT_DIR / f"visual_scene_{digest}.png"
    scene_prompt_base = _visual_prompt(report, "scene")
    scene_ok = False
    for attempt in range(1, 4):
        attempt_prompt = (
            scene_prompt_base
            + f"\nComposition variation {attempt}: use a clearly different camera angle, environment and task moment from the product hero image."
        )
        if ai_configured:
            scene_ok = _generate_ai_visual(attempt_prompt, scene_path)
        else:
            scene_ok = False
        if not scene_ok:
            continue
        similarity = _image_similarity(product_path, scene_path)
        print(f"VISUAL_SIMILARITY_PRODUCT_SCENE: {similarity:.3f} attempt={attempt}", flush=True)
        if similarity < 0.86:
            break
        scene_ok = False

    if not scene_ok:
        if ai_configured and STRICT_AI_IMAGES:
            raise RuntimeError(
                "第2页AI应用场景图生成失败或与封面过于相似。为保证图文质量，已停止Word导出，请重新尝试。"
            )
        scene_ok = _wikimedia_commons_image(report, "scene", scene_path)
        if not scene_ok:
            _fallback_robot_illustration(report, scene_path, scene=True)
    result["scene"] = str(scene_path)

    return result

# ============================================================
# Word 美化与图文排版
# ============================================================
def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_run_font(run, name: str = "Microsoft YaHei", size: float = 10.5, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def _add_title(doc: Document, text: str, size: float = 22) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    _set_run_font(r, size=size, bold=True)


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    _set_run_font(r, size=16 if level == 1 else 13, bold=True)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)


def _add_paragraph(doc: Document, text: Any, bold_prefix: Optional[str] = None, first_line: bool = True) -> None:
    if text is None or str(text).strip() == "":
        return
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.28
    p.paragraph_format.space_after = Pt(3)
    if first_line:
        p.paragraph_format.first_line_indent = Cm(0.74)
    if bold_prefix and str(text).startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        _set_run_font(r1, bold=True)
        r2 = p.add_run(str(text)[len(bold_prefix):])
        _set_run_font(r2)
    else:
        r = p.add_run(str(text))
        _set_run_font(r)


def _add_paragraphs(doc: Document, items: Any, limit: int = 10) -> None:
    if isinstance(items, list):
        for item in items[:limit]:
            _add_paragraph(doc, item)
    elif items:
        _add_paragraph(doc, items)


def _add_bullets(doc: Document, items: Any, limit: int = 10) -> None:
    if not isinstance(items, list):
        if items:
            _add_paragraph(doc, items)
        return
    for item in items[:limit]:
        p = doc.add_paragraph(style=None)
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.3)
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(2)
        if isinstance(item, dict):
            name = item.get("name") or item.get("title") or item.get("dimension") or item.get("function") or ""
            desc = item.get("description") or item.get("implementation") or item.get("ours") or item.get("feedback") or ""
            value = f"• {name}：{desc}" if desc else f"• {name}"
        else:
            value = f"• {item}"
        r = p.add_run(value)
        _set_run_font(r, size=10.2)


def _add_image(doc: Document, path: Optional[str], caption: str, width_cm: float = 15.5) -> None:
    if not path or not Path(path).exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    try:
        r.add_picture(path, width=Cm(width_cm))
    except Exception as exc:
        print("WORD_IMAGE_INSERT_ERROR:", repr(exc), flush=True)
        return
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(3)
    rr = c.add_run(caption)
    _set_run_font(rr, size=9)


def _style_table(table, header_fill: str = "DCE6F1", body_fill: Optional[str] = None) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for ridx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx == 0:
                _set_cell_shading(cell, header_fill)
            elif body_fill:
                _set_cell_shading(cell, body_fill)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    _set_run_font(run, size=9.5, bold=(ridx == 0))


def _add_architecture_table(doc: Document, layers: List[str]) -> None:
    _add_heading(doc, "系统总体架构", 2)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "层级"
    table.rows[0].cells[1].text = "主要职责"
    for item in layers[:6]:
        if "：" in str(item):
            a, b = str(item).split("：", 1)
        else:
            a, b = "模块", str(item)
        cells = table.add_row().cells
        cells[0].text = a
        cells[1].text = b
    _style_table(table)


def _add_flow_table(doc: Document, steps: List[str]) -> None:
    _add_heading(doc, "技术流程", 2)
    steps = [str(s) for s in steps[:8] if s]
    if not steps:
        return
    table = doc.add_table(rows=2, cols=len(steps))
    for i, step in enumerate(steps):
        table.rows[0].cells[i].text = f"步骤{i+1}"
        table.rows[1].cells[i].text = step
    _style_table(table)


def _add_comparison_table(doc: Document, rows: List[Dict[str, Any]]) -> None:
    _add_heading(doc, "差异化对比", 2)
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "比较维度"
    table.rows[0].cells[1].text = "常规方案"
    table.rows[0].cells[2].text = "本项目"
    for item in rows[:8]:
        cells = table.add_row().cells
        cells[0].text = str(item.get("dimension", ""))
        cells[1].text = str(item.get("common", ""))
        cells[2].text = str(item.get("ours", ""))
    _style_table(table)


def _add_function_matrix(doc: Document, rows: List[Dict[str, Any]]) -> None:
    _add_heading(doc, "核心功能矩阵", 2)
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "功能"
    table.rows[0].cells[1].text = "触发方式"
    table.rows[0].cells[2].text = "反馈方式"
    for item in rows[:7]:
        cells = table.add_row().cells
        cells[0].text = str(item.get("function", ""))
        cells[1].text = str(item.get("trigger", ""))
        cells[2].text = str(item.get("feedback", ""))
    _style_table(table)



def _set_cell_text(cell, text: str, size: float = 9.5, bold: bool = False, align: int = WD_ALIGN_PARAGRAPH.CENTER) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    _set_run_font(r, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _diagram_box(doc: Document, title: str, items: List[str], fill: str = "EAF2F8") -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    _set_cell_shading(cell, fill)
    text = title
    if items:
        text += "\n" + " / ".join([str(x) for x in items if str(x).strip()][:5])
    _set_cell_text(cell, text, size=9.8, bold=True)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "10")
        el.set(qn("w:color"), "5B9BD5")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)


def _add_system_block_diagram(doc: Document, diagram: Dict[str, Any]) -> None:
    """报告必备系统框图：单表紧凑版，确保7页物理页数。"""
    _add_heading(doc, "系统框图（必备）", 2)
    stages = [
        ("输入/场景层", diagram.get("inputs", []), "EAF2F8"),
        ("感知处理层", diagram.get("perception", []), "DDEBF7"),
        ("认知决策层", diagram.get("decision", []), "FFF2CC"),
        ("执行服务层", diagram.get("execution", []), "E2F0D9"),
        ("反馈交互层", diagram.get("feedback", []), "FCE4D6"),
    ]
    table = doc.add_table(rows=len(stages), cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (title, items, fill) in enumerate(stages):
        cell = table.cell(i, 0)
        _set_cell_shading(cell, fill)
        body = " / ".join([str(x) for x in items if str(x).strip()][:5])
        arrow = " ↓" if i < len(stages)-1 else ""
        _set_cell_text(cell, f"{title}：{body}{arrow}", size=8.9, bold=True)
    _style_table(table, header_fill="D9EAF7")
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(1)
    rr = cap.add_run("图3  项目系统总体框图")
    _set_run_font(rr, size=8.5)

def _add_hw_sw_block_diagram(doc: Document, diagram: Dict[str, Any]) -> None:
    _add_heading(doc, "软硬件功能框图", 2)
    table = doc.add_table(rows=2, cols=3)
    headers = ["输入/感知端", "主控/软件决策", "执行/输出端"]
    keys = ["input_side", "control_side", "output_side"]
    fills = ["DDEBF7", "FFF2CC", "E2F0D9"]
    for i, h in enumerate(headers):
        _set_cell_shading(table.cell(0, i), fills[i])
        _set_cell_text(table.cell(0, i), h, size=10, bold=True)
        vals = diagram.get(keys[i], []) or []
        _set_cell_text(table.cell(1, i), "\n↓\n".join([str(x) for x in vals[:5]]), size=9.3)
    _style_table(table, header_fill="D9EAF7")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("输入/感知 → 主控决策 → 执行输出，运行状态再回传形成闭环")
    _set_run_font(r, size=9.3, bold=True)


def _add_project_workflow_diagram(doc: Document, steps: List[str]) -> None:
    """项目流程图：动态步骤，Word 原生可编辑。"""
    _add_heading(doc, "项目工作流程图（必备）", 2)
    steps = [str(s) for s in steps if str(s).strip()][:8]
    if not steps:
        return
    cols = len(steps) * 2 - 1
    table = doc.add_table(rows=1, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, step in enumerate(steps):
        c = table.cell(0, i * 2)
        _set_cell_shading(c, "EAF2F8" if i % 2 == 0 else "E2F0D9")
        _set_cell_text(c, step, size=8.5, bold=True)
        if i < len(steps) - 1:
            _set_cell_text(table.cell(0, i * 2 + 1), "→", size=14, bold=True)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = cap.add_run("图4  项目运行与任务闭环流程")
    _set_run_font(rr, size=9)


def _add_landing_roadmap(doc: Document, steps: List[str]) -> None:
    _add_heading(doc, "项目落地路线图", 2)
    steps = [str(s) for s in steps if str(s).strip()][:6]
    if not steps:
        return
    table = doc.add_table(rows=len(steps), cols=2)
    for i, step in enumerate(steps):
        _set_cell_shading(table.cell(i, 0), "D9EAF7")
        _set_cell_text(table.cell(i, 0), f"阶段 {i+1}", size=9.4, bold=True)
        clean = re.sub(r"^阶段\d+[:：]", "", step)
        _set_cell_text(table.cell(i, 1), clean, size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT)
    _style_table(table, header_fill="D9EAF7")


def _add_revision_notes(doc: Document, content: Dict[str, Any]) -> None:
    notes = content.get("user_requested_changes") or []
    if not notes:
        return
    _add_heading(doc, "用户定制补充", 2)
    _add_bullets(doc, notes, 8)

def _prepare_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.45)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.3)
    return doc


def _render_cover(doc: Document, report: Dict[str, Any], images: Dict[str, str]) -> None:
    doc.add_paragraph("\n")
    _add_title(doc, report.get("project_title", "智能机器人竞赛项目"), 22)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("智能机器人创意竞赛 · 项目设计报告")
    _set_run_font(r, size=13)
    _add_image(doc, images.get("product"), "图1  项目产品概念效果图", 15.8)

    content = report.get("pages", {}).get("page_1", {}).get("content", {})
    _add_paragraph(doc, content.get("design_statement", ""), first_line=False)

    selected = report.get("selected_candidate", {})
    table = doc.add_table(rows=4, cols=2)
    pairs = [
        ("项目定位", selected.get("positioning", "")),
        ("核心技术", selected.get("core_tech", "")),
        ("差异化", selected.get("differentiator", "")),
        ("竞赛竞争力", f"{selected.get('scores', {}).get('competitiveness_score', '')} / 100"),
    ]
    for row, (k, v) in zip(table.rows, pairs):
        row.cells[0].text = k
        row.cells[1].text = str(v)
    _style_table(table, header_fill="E8EEF7")


def _render_page(doc: Document, report: Dict[str, Any], page_no: int, images: Dict[str, str]) -> None:
    page = report["pages"][f"page_{page_no}"]
    content = page.get("content", {})
    _add_heading(doc, f"第{page_no}页  {page.get('title', '')}", 1)

    if page_no == 2:
        _add_image(doc, images.get("scene"), "图2  项目典型应用场景效果图", 11.5)
        _add_heading(doc, "设计背景", 2)
        _add_paragraphs(doc, content.get("background_paragraphs") or content.get("background", []), 2)
        _add_heading(doc, "用户痛点", 2)
        _add_bullets(doc, content.get("pain_points", []), 3)
        _add_heading(doc, "用户需求", 2)
        _add_bullets(doc, content.get("user_needs", []), 3)
        _add_heading(doc, "设计目标", 2)
        _add_bullets(doc, content.get("design_goals", []), 3)

    elif page_no == 3:
        _add_paragraph(doc, content.get("overview", ""))
        diagram = content.get("system_block_diagram") or {
            "inputs": ["用户指令", "场景/对象状态", "传感器数据"],
            "perception": ["多源感知", "数据预处理"],
            "decision": ["状态识别", "任务规划", "安全判断"],
            "execution": content.get("core_modules", [])[:4],
            "feedback": ["执行确认", "异常提醒", "远程同步"],
        }
        _add_system_block_diagram(doc, diagram)
        _add_architecture_table(doc, content.get("system_layers", []))
        _add_heading(doc, "核心模块", 2)
        _add_bullets(doc, content.get("core_modules", []), 4)
        _add_heading(doc, "本体结构思路", 2)
        _add_paragraph(doc, content.get("mechanical_concept", ""))
        _add_heading(doc, "结构设计原则", 2)
        _add_bullets(doc, content.get("design_principles", []), 3)

    elif page_no == 4:
        _add_hw_sw_block_diagram(doc, content.get("hardware_software_block", {}))
        _add_heading(doc, "硬件设计", 2)
        _add_bullets(doc, content.get("hardware", []), 8)
        _add_heading(doc, "软件功能", 2)
        _add_bullets(doc, content.get("software", []), 8)
        _add_function_matrix(doc, content.get("functional_matrix", []))
        _add_heading(doc, "交互逻辑", 2)
        _add_paragraph(doc, content.get("interaction_logic", ""), first_line=False)

    elif page_no == 5:
        _add_paragraph(doc, content.get("technical_summary", ""))
        _add_project_workflow_diagram(doc, content.get("project_workflow") or content.get("technical_route", []))
        _add_flow_table(doc, content.get("technical_route", []))
        _add_heading(doc, "关键技术", 2)
        _add_bullets(doc, content.get("key_technologies", []), 6)
        _add_heading(doc, "量化指标", 2)
        _add_bullets(doc, content.get("engineering_metrics", []), 6)
        _add_heading(doc, "风险控制", 2)
        _add_bullets(doc, content.get("risk_control", []), 6)

    elif page_no == 6:
        _add_paragraph(doc, content.get("innovation_intro", ""))
        _add_heading(doc, "创新点", 2)
        _add_bullets(doc, content.get("innovation_points", []), 7)
        _add_comparison_table(doc, content.get("comparison", []))
        diff = content.get("differentiation_analysis", {})
        note = diff.get("current_boundary") if isinstance(diff, dict) else None
        if note:
            _add_heading(doc, "知识库说明", 2)
            _add_paragraph(doc, note)

    elif page_no == 7:
        _add_paragraphs(doc, content.get("prospect_paragraphs", []), 4)
        _add_heading(doc, "典型应用场景", 2)
        _add_bullets(doc, content.get("application_scenarios", []), 6)
        _add_landing_roadmap(doc, content.get("deployment_paths", []))
        _add_heading(doc, "社会价值", 2)
        _add_bullets(doc, content.get("social_value", []), 6)
        _add_heading(doc, "后续迭代", 2)
        _add_bullets(doc, content.get("future_iterations", []), 6)

    _add_revision_notes(doc, content)

def export_report_to_word(report_json: Dict[str, Any], session_key: str, fast_mode: bool = False) -> Dict[str, str]:
    images = generate_project_images(report_json, session_key, fast_mode=fast_mode)
    required_images = ["product", "scene"]
    for image_key in required_images:
        image_path = Path(images.get(image_key, ""))
        if not image_path.is_file() or image_path.stat().st_size < 2048:
            raise RuntimeError(f"Required report illustration missing or too small: {image_key}")

    doc = _prepare_document()
    _render_cover(doc, report_json, images)

    for page_no in range(2, 8):
        doc.add_page_break()
        _render_page(doc, report_json, page_no, images)

    digest = hashlib.sha1(
        f"{session_key}-{report_json.get('project_title','')}-{now_iso()}".encode("utf-8")
    ).hexdigest()[:12]

    # ASCII filename avoids Chaoxing/browser filename parsing problems.
    filename = f"robot_competition_report_{digest}.docx"
    path = EXPORT_DIR / filename
    doc.save(path)

    # Verify the .docx actually contains embedded raster images before reporting success.
    try:
        with zipfile.ZipFile(path, "r") as archive:
            embedded_images = [
                name for name in archive.namelist()
                if name.startswith("word/media/") and not name.endswith("/")
            ]
    except Exception as exc:
        raise RuntimeError(f"Unable to verify generated Word document: {type(exc).__name__}") from exc

    if len(embedded_images) < 2:
        raise RuntimeError(f"Word export has only {len(embedded_images)} embedded images; expected at least 2.")

    file_size = path.stat().st_size
    project_title = str(report_json.get("project_title", ""))
    print(
        "WORD_EXPORT_VERIFIED:",
        {
            "filename": filename,
            "project_title": project_title,
            "embedded_images": len(embedded_images),
            "size_bytes": file_size,
            "fast_mode": fast_mode,
        },
        flush=True,
    )

    return {
        "filename": filename,
        "file_path": str(path),
        "download_url": f"{PUBLIC_BASE_URL}/api/v1/robot-competition/download/{filename}",
        "visual_mode": "offline_stylized_illustrations_plus_native_diagrams" if fast_mode else ("ai_generated_strict" if (POLLINATIONS_API_KEY or IMAGE_API_URL) else "offline_fallback"),
        "project_title": project_title,
        "embedded_images": len(embedded_images),
        "size_bytes": file_size,
    }

# ============================================================
# 意图识别与会话流程
# ============================================================
def detect_page_number(message: str) -> Optional[int]:
    text = normalize_text(message)
    cn = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7}
    for pattern in [r"第([1-7])页", r"page([1-7])", r"页([1-7])"]:
        m = re.search(pattern, text, re.I)
        if m:
            return int(m.group(1))
    m = re.search(r"第([一二三四五六七])页", text)
    if m:
        return cn.get(m.group(1))
    # 用户未写页码时，根据明确模块自动路由到对应页面
    if contains_any(text, ["系统框图", "系统架构", "总体结构", "整体结构"]):
        return 3
    if contains_any(text, ["软硬件", "硬件设计", "软件功能", "功能框图"]):
        return 4
    if contains_any(text, ["项目流程图", "工作流程图", "技术路线", "流程图"]):
        return 5
    if contains_any(text, ["创新点", "差异化", "对比图"]):
        return 6
    if contains_any(text, ["行业前景", "应用前景", "落地路线", "落地路径"]):
        return 7
    if contains_any(text, ["设计背景", "用户需求", "用户痛点"]):
        return 2
    if contains_any(text, ["封面", "项目题目", "标题"]):
        return 1
    return None


def _looks_like_new_project_idea(message: str) -> bool:
    """判断用户是在描述一个新的机器人创意，而不是在修改已生成报告。

    规则故意偏保守：只要出现明显的“我要做/设计一个机器人”语义，就优先视为新项目。
    这样即使同一会话里已经有旧 report_json，也不会把新创意误判成“修改报告”。
    """
    text = normalize_text(message)
    strong_starts = [
        "我想设计一个", "我想做一个", "我想设计", "我想做", "我准备设计", "我准备做",
        "设计一个", "做一个", "开发一个", "制作一个", "构思一个", "我的创意是", "机器人创意",
    ]
    robot_terms = ["机器人", "智能小车", "机械臂", "无人车", "巡检车", "陪护", "智能设备"]
    if any(text.startswith(x) for x in strong_starts) and contains_any(text, robot_terms):
        return True
    # 较长、包含功能描述的机器人需求，也按新创意处理
    function_terms = ["可以", "用于", "主要", "实现", "能够", "自动", "识别", "检测", "提醒", "陪伴", "清理", "通知"]
    return len(text) >= 28 and contains_any(text, robot_terms) and sum(w in text for w in function_terms) >= 2


def _is_explicit_revision_request(message: str, session: Dict[str, Any]) -> bool:
    """报告生成后才允许进入修改流程，而且必须有明确修改意图。"""
    if not session.get("report_json"):
        return False
    text = normalize_text(message)
    revision_words = ["修改", "调整", "更改", "改成", "修改为", "补充", "增加", "新增", "删除", "删掉", "替换", "优化", "重写"]
    if not contains_any(text, revision_words):
        return False
    # 新项目描述永远优先，不误判成修改
    if _looks_like_new_project_idea(message):
        return False
    # 明确页码或报告模块时才算修改。
    return bool(
        detect_page_number(message)
        or contains_any(text, ["系统框图", "流程图", "工作流程", "软硬件", "创新点", "差异化", "应用前景", "落地路线", "设计背景", "用户需求", "封面", "标题"])
    )


def detect_intent(message: str, session: Dict[str, Any]) -> str:
    text = normalize_text(message)

    # 0. 无论旧会话处于什么阶段，只要用户明确给出一个新的机器人创意，就重新开始项目。
    if _looks_like_new_project_idea(message):
        return "create_project"

    # 1. 报告查询/修改只允许发生在报告已经生成后。
    if contains_any(text, ["查看报告结构", "查看结构化数据", "查看json"]):
        return "view_report_json"
    if _is_explicit_revision_request(message, session):
        return "modify_page_request"

    # 2. 下载优先于“重新生成题目”，避免“重新导出Word”之类被误判。
    if contains_any(text, ["下载", "导出", "生成word", "word报告", "word文档"]) and contains_any(text, ["报告", "word", "文档", "方案书"]):
        return "word_export_pending"

    # 3. 报告生成阶段。
    if contains_any(text, ["生成报告", "写报告", "生成方案书"]):
        return "generate_report"

    # 4. 选题阶段。
    if re.fullmatch(r"[123]", text) or contains_any(text, ["选1", "选2", "选3", "第一个", "第二个", "第三个"]):
        return "select_title"
    if contains_any(text, ["重新生成", "重生成", "再生成", "换一批", "题目不满意", "换题"]):
        return "regenerate_titles"

    # 5. 首次消息就是新项目；在候选题阶段可以继续补充创意要求。
    if "raw_idea" not in session:
        return "create_project"
    if session.get("stage") in {"candidates_ready", "title_confirmed"} and not session.get("report_json"):
        return "supplement_idea"

    # 6. 报告已经生成后，普通闲聊/新需求不擅自改报告。
    if session.get("report_json"):
        return "post_report_chat"
    return "supplement_idea"

def parse_selection(message: str) -> Optional[int]:
    text = normalize_text(message)
    if text == "1" or "选1" in text or "第一个" in text:
        return 1
    if text == "2" or "选2" in text or "第二个" in text:
        return 2
    if text == "3" or "选3" in text or "第三个" in text:
        return 3
    return None


def response(**kwargs: Any) -> ChatResponse:
    kwargs.setdefault("payload_json", "")
    kwargs.setdefault("download_url", "")
    kwargs.setdefault("data", {})
    kwargs.setdefault("files", [])
    kwargs.setdefault("suggested_actions", [])
    kwargs.setdefault("updated_at", now_iso())
    return ChatResponse(**kwargs)



def _apply_local_revision(report_json: Dict[str, Any], page_no: int, request: str) -> None:
    key = f"page_{page_no}"
    page = report_json.get("pages", {}).get(key)
    if not isinstance(page, dict):
        return
    content = page.setdefault("content", {})
    notes = content.setdefault("user_requested_changes", [])
    if request not in notes:
        notes.append(request)

    # 常见修改可直接落地，不依赖外部LLM。
    text = request.strip()
    if page_no == 1:
        m = re.search(r"(?:题目|标题).{0,8}(?:改成|修改为|更改为|换成)[：:\s]*[“\"']?(.+?)[”\"']?$", text)
        if m and m.group(1).strip():
            new_title = m.group(1).strip().strip("。")
            report_json["project_title"] = new_title
            page["title"] = new_title
            if isinstance(report_json.get("selected_candidate"), dict):
                report_json["selected_candidate"]["title"] = new_title
    elif page_no == 3 and contains_any(normalize_text(text), ["知识库", "云端", "大模型"]):
        diag = content.setdefault("system_block_diagram", {})
        decision = diag.setdefault("decision", [])
        addition = "知识库/智能推理"
        if addition not in decision:
            decision.append(addition)
    elif page_no == 5 and contains_any(normalize_text(text), ["测试", "验证", "评价"]):
        flow = content.setdefault("project_workflow", [])
        if "测试验证与指标评价" not in flow:
            flow.append("测试验证与指标评价")

def handle_chat(req: ChatRequest) -> ChatResponse:
    session_key = build_session_key(req)
    session = load_session(session_key)
    user_message = clean_message(req.message)
    intent = detect_intent(user_message, session)

    print("DEBUG_MESSAGE:", user_message, "INTENT:", intent, flush=True)
    print("DEBUG_SESSION:", session_key, "SESSION_KEYS:", list(session.keys()), flush=True)
    print("DEBUG_SESSION_STATE:", {"session_key": session_key, "stage": session.get("stage"), "has_report_json": bool(session.get("report_json")), "has_selected_title": bool(session.get("selected_title"))}, flush=True)

    if intent in {"create_project", "supplement_idea", "regenerate_titles"}:
        if intent == "create_project":
            raw_idea = user_message
            prefix = "已收到你的机器人创意，正在根据本次输入动态生成3个候选题目。"
        elif intent == "supplement_idea":
            raw_idea = (session.get("raw_idea", "") + "\n补充要求：" + user_message).strip()
            prefix = "已把补充要求合并到原始创意，下面重新生成候选题目。"
        else:
            raw_idea = session.get("raw_idea", user_message)
            # 增加扰动，保证“重新生成”能产生不同命名倾向
            raw_idea = raw_idea + f"\n重新生成批次：{now_iso()}"
            prefix = "已根据当前创意重新生成一批候选题目。"

        fields = extract_idea_fields(raw_idea)
        knowledge = search_competition_knowledge(fields)
        candidates = generate_candidates(fields, knowledge)

        session.update({
            "raw_idea": raw_idea,
            "fields": fields,
            "knowledge": knowledge,
            "candidates": candidates,
            "selected_title": None,
            "report_json": None,
            "stage": "candidates_ready",
        })
        save_session(session_key, session)

        payload = {"fields": fields, "candidates": candidates, "knowledge": knowledge}
        return response(
            success=True,
            stage="candidates_ready",
            intent=intent,
            message=format_candidate_message(fields, candidates, prefix, knowledge),
            payload_json=json.dumps(payload, ensure_ascii=False, indent=2),
            data=payload,
            suggested_actions=["输入1", "输入2", "输入3", "重新生成", "补充要求"],
        )

    if intent == "select_title":
        selected_id = parse_selection(user_message)
        candidates = session.get("candidates", [])
        selected = next((c for c in candidates if c.get("id") == selected_id), None)
        if not selected:
            return response(
                success=False,
                stage="selection_failed",
                intent=intent,
                message="没有找到对应候选题目。请先输入一个完整机器人创意，然后选择1、2或3。",
                suggested_actions=["输入1", "输入2", "输入3"],
            )

        session["selected_title"] = selected
        session["stage"] = "title_confirmed"
        save_session(session_key, session)
        payload = {"selected_title": selected, "current_stage": "title_confirmed"}
        return response(
            success=True,
            stage="title_confirmed",
            intent=intent,
            message=(
                f"已确认最终竞赛题目：\n\n{selected['title']}\n\n"
                f"竞赛竞争力评分：{selected['scores']['competitiveness_score']} / 100\n"
                f"核心技术：{selected['core_tech']}\n\n"
                "请输入“生成报告”，系统将围绕本项目生成固定7页结构、动态内容的竞赛报告。"
            ),
            payload_json=json.dumps(payload, ensure_ascii=False, indent=2),
            data=payload,
            suggested_actions=["生成报告", "重新生成题目"],
        )

    if intent == "generate_report":
        selected = session.get("selected_title")
        fields = session.get("fields")
        knowledge = session.get("knowledge") or {}
        if not selected or not fields:
            return response(
                success=False,
                stage="report_failed",
                intent=intent,
                message="还没有确认最终题目。请先输入机器人创意并选择1、2或3。",
                suggested_actions=["输入1", "输入2", "输入3"],
            )

        report_json = build_report_json(fields, selected, knowledge)
        session["report_json"] = report_json
        session["report_revision"] = int(session.get("report_revision", 0)) + 1
        session["stage"] = "report_json_ready"
        session.pop("word_file", None)
        save_session(session_key, session)

        # Create the Word file in the same request as report generation.
        # This avoids relying on a separate Chaoxing "download report" node to fire.
        try:
            export_info = export_report_to_word(report_json, session_key, fast_mode=True)
            session["word_file"] = export_info
            session["stage"] = "word_ready"
            save_session(session_key, session)
            download_url = export_info["download_url"]
            report_message = (
                format_report_summary(report_json)
                + "\n\nWord报告已自动生成（7页）。文档已验证包含"
                + str(export_info.get("embedded_images", 0))
                + "张嵌入图片，以及系统框图、软硬件功能框图和项目流程图。"
                + "\n\nWord下载地址（请点击或复制完整地址）：\n"
                + download_url
                + "\n\n对应题目："
                + str(report_json.get("project_title", ""))
            )
            return response(
                success=True,
                stage="word_ready",
                intent=intent,
                message=report_message,
                payload_json=json.dumps(report_json, ensure_ascii=False, indent=2),
                download_url=download_url,
                data={"report_json": report_json, "word_file": export_info, "revision": session["report_revision"]},
                files=[{"name": export_info["filename"], "url": download_url, "type": "docx"}],
                suggested_actions=["下载报告", "查看报告结构数据", "修改指定页面后重新导出"],
            )
        except Exception as exc:
            print("DEBUG_AUTO_WORD_ERROR:", repr(exc), flush=True)
            session["stage"] = "report_json_ready"
            save_session(session_key, session)
            return response(
                success=True,
                stage="report_json_ready",
                intent=intent,
                message=(
                    format_report_summary(report_json)
                    + "\\n\\nWord文件暂未生成（" + type(exc).__name__ + "）。"
                    + "请继续输入“下载报告”重试。"
                ),
                payload_json=json.dumps(report_json, ensure_ascii=False, indent=2),
                data={"report_json": report_json, "revision": session["report_revision"], "word_export_error": type(exc).__name__},
                suggested_actions=["下载报告", "查看报告结构数据"],
            )

    if intent == "view_report_json":
        report_json = session.get("report_json")
        if not report_json:
            return response(
                success=False,
                stage="report_not_ready",
                intent=intent,
                message="当前还没有报告数据。请先选择题目并输入“生成报告”。",
            )
        return response(
            success=True,
            stage="report_json_ready",
            intent=intent,
            message="报告结构化JSON已返回到 payload_json 字段。",
            payload_json=json.dumps(report_json, ensure_ascii=False, indent=2),
            data={"report_json": report_json},
        )

    if intent == "modify_page_request":
        report_json = session.get("report_json")
        if not report_json:
            return response(
                success=False,
                stage="report_not_ready",
                intent=intent,
                message="当前还没有报告。请先生成报告。",
            )
        page_no = detect_page_number(user_message)
        if not page_no:
            return response(
                success=False,
                stage="modify_need_target",
                intent=intent,
                message="可以修改。请告诉我要修改第1至第7页中的哪一页，例如：‘修改第3页，在系统框图中增加云端知识库模块’。",
                suggested_actions=["修改第3页系统框图", "修改第5页项目流程图", "修改第6页创新点"],
            )

        key = f"page_{page_no}"
        page = report_json["pages"].get(key, {})
        page.setdefault("revision_notes", [])
        page["revision_notes"].append(user_message)

        # 先执行本地可确定的修改，保证没有外部LLM时也确实发生变化。
        _apply_local_revision(report_json, page_no, user_message)
        page = report_json["pages"].get(key, page)

        # 如配置LLM，再按自然语言做更深入的重写。
        ai = call_llm_json(
            "你是竞赛报告编辑专家。只返回JSON对象，保持page_no/module/title/content结构。严格依据用户修改要求重写该页；系统框图、流程图等结构字段必须保留为可渲染的列表/字典。",
            json.dumps({"page": page, "request": user_message, "project_title": report_json.get("project_title")}, ensure_ascii=False),
        )
        if ai and isinstance(ai, dict):
            ai.setdefault("page_no", page_no)
            ai.setdefault("module", page.get("module"))
            ai.setdefault("title", page.get("title"))
            ai.setdefault("content", page.get("content"))
            # 保留本地修改记录
            ai["revision_notes"] = page.get("revision_notes", [])
            local_notes = page.get("content", {}).get("user_requested_changes", [])
            ai.setdefault("content", {}).setdefault("user_requested_changes", local_notes)
            report_json["pages"][key] = ai

        session["report_json"] = report_json
        session["report_revision"] = int(session.get("report_revision", 1)) + 1
        save_session(session_key, session)
        return response(
            success=True,
            stage="report_json_ready",
            intent=intent,
            message=(
                f"已完成第{page_no}页修改。当前版本：{session['report_revision']}。\n"
                "你可以继续提出修改要求，或输入‘下载报告’重新导出最新版Word。"
            ),
            payload_json=json.dumps(report_json, ensure_ascii=False, indent=2),
            data={"report_json": report_json, "revision": session["report_revision"]},
            suggested_actions=["下载报告", "继续修改第3页", "修改第5页流程图"],
        )

    if intent == "word_export_pending":
        report_json = session.get("report_json")
        if not report_json:
            return response(
                success=False,
                stage="word_export_failed",
                intent=intent,
                message="当前还没有报告数据。请先完成选题并输入“生成报告”。",
                suggested_actions=["生成报告"],
            )

        try:
            export_info = export_report_to_word(report_json, session_key, fast_mode=True)
            print("DEBUG_WORD_SUCCESS:", export_info, flush=True)
        except Exception as exc:
            print("DEBUG_WORD_ERROR:", repr(exc), flush=True)
            return response(
                success=False,
                stage="word_export_failed",
                intent=intent,
                message=f"Word生成失败：{type(exc).__name__}。请查看Render日志。",
                payload_json=json.dumps(report_json, ensure_ascii=False, indent=2),
                data={"error": str(exc)},
            )

        session["word_file"] = export_info
        session["stage"] = "word_ready"
        save_session(session_key, session)

        return response(
            success=True,
            stage="word_ready",
            intent=intent,
            message=(
                "Word报告已经生成完成（7页，已嵌入"
                + str(export_info.get("embedded_images", 0))
                + "张图片）。\n\n"
                "对应题目：" + str(report_json.get("project_title", ""))
                + "\n\nWord下载地址（请点击或复制完整地址）：\n"
                + export_info["download_url"]
            ),
            payload_json=json.dumps(report_json, ensure_ascii=False, indent=2),
            download_url=export_info["download_url"],
            data={"word_file": export_info, "report_json": report_json},
            files=[{"name": export_info["filename"], "url": export_info["download_url"], "type": "docx"}],
            suggested_actions=["下载报告", "修改指定页面后重新导出"],
        )

    if intent == "post_report_chat":
        return response(
            success=True,
            stage="report_json_ready",
            intent=intent,
            message=(
                "当前报告已经生成。\n\n"
                "如果你对报告不满意，请明确告诉我要修改哪一页或哪个模块，例如：\n"
                "• 修改第3页，在系统框图中增加云端知识库模块\n"
                "• 优化第5页项目工作流程图，增加测试验证环节\n"
                "• 修改第6页，突出本项目相对传统方案的创新点\n\n"
                "如果你要做一个全新的机器人项目，直接输入新的完整创意即可，我会自动重新生成3个候选题目。"
            ),
            payload_json=json.dumps(session.get("report_json", {}), ensure_ascii=False, indent=2),
            data={"report_json": session.get("report_json", {}), "revision": session.get("report_revision", 1)},
            suggested_actions=["下载报告", "修改第3页系统框图", "优化第5页流程图"],
        )

    return response(
        success=False,
        stage="unknown",
        intent="unknown",
        message="暂时无法识别该操作。你可以输入机器人创意、1/2/3、生成报告、修改指定页面或下载报告。",
    )


# ============================================================
# API 路由
# ============================================================
@app.get("/")
def root() -> Dict[str, Any]:
    return {
        "service": "robot-competition-orchestrator",
        "version": APP_VERSION,
        "status": "ok",
        "llm_enabled": llm_available(),
        "knowledge_base_enabled": bool(KNOWLEDGE_API_URL),
    }


@app.get("/health")
def health() -> Dict[str, Any]:
    return {"ok": True, "version": APP_VERSION, "time": now_iso()}


@app.get("/api/v1/diagnostics/knowledge")
def diagnose_knowledge() -> Dict[str, Any]:
    """Safely verify the configured knowledge API end-to-end without exposing secrets or document contents."""
    if not KNOWLEDGE_API_URL:
        return {
            "ok": False,
            "configured": False,
            "knowledge_base_used": False,
            "source_type": "not_configured",
            "returned_results": 0,
            "message": "KNOWLEDGE_API_URL is not configured.",
        }

    fields = {
        "keywords": ["智能小车", "机器人竞赛", "获奖项目"],
        "core_functions": ["自主导航", "避障"],
        "scenarios": ["校园"],
    }
    result = search_competition_knowledge(fields, top_k=3)
    cases = result.get("similar_award_cases") or []
    used = bool(result.get("knowledge_base_used"))
    return {
        "ok": used,
        "configured": True,
        "knowledge_base_used": used,
        "source_type": result.get("source_type", "unknown"),
        "returned_results": len(cases) if used else 0,
        "sample_fallback": result.get("source_type") == "sample_case_library",
        "message": "Live knowledge search succeeded." if used else "The service is configured, but live search did not return usable records; inspect Render logs for status/auth errors.",
    }


@app.post("/api/v1/robot-competition/chat", response_model=ChatResponse)
def chat(req: ChatRequest) -> ChatResponse:
    return handle_chat(req)


@app.get("/api/v1/robot-competition/download/{filename}")
def download(filename: str):
    # Some workflow renderers incorrectly include the Markdown closing ")" in the URL.
    # Trim only trailing punctuation; keep the .docx extension intact.
    cleaned_filename = filename.rstrip("),;]}>""'")
    safe = Path(cleaned_filename).name
    file_path = EXPORT_DIR / safe
    exists = file_path.exists() and file_path.is_file()
    print(
        "DEBUG_DOWNLOAD_REQUEST:",
        {"filename": safe, "exists": exists, "size_bytes": file_path.stat().st_size if exists else 0},
        flush=True,
    )
    if not exists:
        raise HTTPException(status_code=404, detail="File not found")
    return FileResponse(
        path=str(file_path),
        filename=safe,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
