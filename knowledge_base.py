from __future__ import annotations

import asyncio
import hashlib
import hmac
import logging
import io
import json
import os
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, Response
from pydantic import BaseModel, Field
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
from sqlalchemy import (
    DateTime, ForeignKey, Integer, LargeBinary, String, Text, create_engine,
    select, func, event, inspect as sa_inspect, text
)
from sqlalchemy.orm import DeclarativeBase, Mapped, mapped_column, relationship, sessionmaker

try:
    import boto3
    from botocore.config import Config as BotoConfig
except ImportError:
    boto3 = None
    BotoConfig = None

logging.basicConfig(
    level=os.getenv("LOG_LEVEL", "INFO").upper(),
    format="%(asctime)s %(levelname)s %(name)s: %(message)s",
)
logger = logging.getLogger("robot_competition_knowledge_base")

APP_VERSION = "0.2.0"
MAX_FILE_BYTES = int(os.getenv("MAX_FILE_BYTES", str(25 * 1024 * 1024)))
MAX_UPLOADS = int(os.getenv("MAX_UPLOADS", "30"))
COOKIE_NAME = "robot_kb_session"
SESSION_SECRET = os.getenv("SESSION_SECRET", "")
KB_API_KEY = os.getenv("KB_API_KEY", "")
COOKIE_SECURE = os.getenv("COOKIE_SECURE", "true").lower() not in {"0", "false", "no"}

DATABASE_URL = os.getenv("DATABASE_URL", "").strip()

# Preferred storage: Neon Object Storage (S3-compatible, no separate card-based R2 setup).
# Legacy Cloudflare R2 variables remain supported for backward compatibility.
AWS_ENDPOINT_URL_S3 = os.getenv("AWS_ENDPOINT_URL_S3", "").strip()
AWS_ACCESS_KEY_ID = os.getenv("AWS_ACCESS_KEY_ID", "").strip()
AWS_SECRET_ACCESS_KEY = os.getenv("AWS_SECRET_ACCESS_KEY", "").strip()
AWS_REGION = os.getenv("AWS_REGION", "us-east-2").strip() or "us-east-2"
S3_BUCKET_NAME = os.getenv("S3_BUCKET_NAME", "").strip()

R2_ACCOUNT_ID = os.getenv("R2_ACCOUNT_ID", "").strip()
R2_ACCESS_KEY_ID = os.getenv("R2_ACCESS_KEY_ID", "").strip()
R2_SECRET_ACCESS_KEY = os.getenv("R2_SECRET_ACCESS_KEY", "").strip()
R2_BUCKET_NAME = os.getenv("R2_BUCKET_NAME", "").strip()
R2_ENDPOINT_URL = os.getenv("R2_ENDPOINT_URL", "").strip() or (f"https://{R2_ACCOUNT_ID}.r2.cloudflarestorage.com" if R2_ACCOUNT_ID else "")

if AWS_ENDPOINT_URL_S3:
    OBJECT_STORAGE_ENDPOINT_URL = AWS_ENDPOINT_URL_S3
    OBJECT_STORAGE_ACCESS_KEY_ID = AWS_ACCESS_KEY_ID
    OBJECT_STORAGE_SECRET_ACCESS_KEY = AWS_SECRET_ACCESS_KEY
    OBJECT_STORAGE_BUCKET_NAME = S3_BUCKET_NAME
    OBJECT_STORAGE_REGION = AWS_REGION
    OBJECT_STORAGE_PROVIDER = "neon_object_storage"
elif R2_ENDPOINT_URL:
    OBJECT_STORAGE_ENDPOINT_URL = R2_ENDPOINT_URL
    OBJECT_STORAGE_ACCESS_KEY_ID = R2_ACCESS_KEY_ID
    OBJECT_STORAGE_SECRET_ACCESS_KEY = R2_SECRET_ACCESS_KEY
    OBJECT_STORAGE_BUCKET_NAME = R2_BUCKET_NAME
    OBJECT_STORAGE_REGION = "auto"
    OBJECT_STORAGE_PROVIDER = "cloudflare_r2"
else:
    OBJECT_STORAGE_ENDPOINT_URL = ""
    OBJECT_STORAGE_ACCESS_KEY_ID = ""
    OBJECT_STORAGE_SECRET_ACCESS_KEY = ""
    OBJECT_STORAGE_BUCKET_NAME = ""
    OBJECT_STORAGE_REGION = "us-east-2"
    OBJECT_STORAGE_PROVIDER = "database_blob"

OBJECT_STORAGE_ENABLED = bool(
    boto3 and OBJECT_STORAGE_ENDPOINT_URL and OBJECT_STORAGE_ACCESS_KEY_ID
    and OBJECT_STORAGE_SECRET_ACCESS_KEY and OBJECT_STORAGE_BUCKET_NAME
)
OBJECT_STORAGE_CLIENT = boto3.client(
    "s3",
    endpoint_url=OBJECT_STORAGE_ENDPOINT_URL,
    aws_access_key_id=OBJECT_STORAGE_ACCESS_KEY_ID,
    aws_secret_access_key=OBJECT_STORAGE_SECRET_ACCESS_KEY,
    region_name=OBJECT_STORAGE_REGION,
    config=BotoConfig(signature_version="s3v4", s3={"addressing_style": "path"}),
) if OBJECT_STORAGE_ENABLED else None
R2_ENABLED = OBJECT_STORAGE_ENABLED and OBJECT_STORAGE_PROVIDER == "cloudflare_r2"
R2_CLIENT = OBJECT_STORAGE_CLIENT if R2_ENABLED else None
if DATABASE_URL:
    if DATABASE_URL.startswith("postgres://"):
        DATABASE_URL = "postgresql+psycopg://" + DATABASE_URL[len("postgres://"):]
    elif DATABASE_URL.startswith("postgresql://") and "+psycopg" not in DATABASE_URL:
        DATABASE_URL = "postgresql+psycopg://" + DATABASE_URL[len("postgresql://"):]
    engine = create_engine(
        DATABASE_URL,
        pool_pre_ping=True,
        connect_args={
            "connect_timeout": 8,
            "options": "-c statement_timeout=15000 -c lock_timeout=5000",
        },
    )
else:
    db_path = Path(os.getenv("KB_DB_PATH", "/tmp/robot_kb.sqlite3"))
    db_path.parent.mkdir(parents=True, exist_ok=True)
    engine = create_engine(
        "sqlite:///" + str(db_path),
        connect_args={"check_same_thread": False},
    )
    @event.listens_for(engine, "connect")
    def _sqlite_foreign_keys(dbapi_connection, connection_record):
        cursor = dbapi_connection.cursor()
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()

SessionLocal = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)


class Base(DeclarativeBase):
    pass


class KnowledgeDocument(Base):
    __tablename__ = "knowledge_documents"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    original_name: Mapped[str] = mapped_column(String(255), nullable=False)
    sha256: Mapped[str] = mapped_column(String(64), unique=True, index=True, nullable=False)
    content_type: Mapped[str] = mapped_column(String(160), default="application/octet-stream")
    file_data: Mapped[bytes | None] = mapped_column(LargeBinary, nullable=True)
    object_key: Mapped[str | None] = mapped_column(String(1024), nullable=True)
    extracted_text: Mapped[str] = mapped_column(Text, default="")
    status: Mapped[str] = mapped_column(String(40), default="已入库")
    uploaded_by: Mapped[str] = mapped_column(String(80), default="unknown")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    chunks: Mapped[list["KnowledgeChunk"]] = relationship(
        back_populates="document", cascade="all, delete-orphan"
    )


class KnowledgeChunk(Base):
    __tablename__ = "knowledge_chunks"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    document_id: Mapped[int] = mapped_column(ForeignKey("knowledge_documents.id", ondelete="CASCADE"), index=True)
    chunk_index: Mapped[int] = mapped_column(Integer, default=0)
    content: Mapped[str] = mapped_column(Text, nullable=False)
    document: Mapped[KnowledgeDocument] = relationship(back_populates="chunks")


DATABASE_READY = False
DATABASE_INIT_ERROR = ""
FILE_DATA_NOT_NULL = True


def initialize_database() -> bool:
    """Initialize SQLite locally; reuse the already-migrated Render PostgreSQL schema."""
    global DATABASE_READY, DATABASE_INIT_ERROR, FILE_DATA_NOT_NULL
    try:
        logger.info("Starting database initialization (dialect=%s).", engine.dialect.name)
        if engine.dialect.name == "postgresql":
            # Render's existing production database already contains both knowledge-base
            # tables and the object_key migration. Re-running SQLAlchemy create_all()
            # triggers PostgreSQL catalog introspection that can be blocked by DDL locks.
            # Reuse the known deployed schema and let real queries surface any schema issue.
            FILE_DATA_NOT_NULL = False
            with engine.connect() as connection:
                connection.execute(text("SELECT 1"))
            logger.info("PostgreSQL connection verified; reusing existing schema.")
        else:
            Base.metadata.create_all(engine)
            with engine.begin() as connection:
                columns = {
                    column["name"]: column
                    for column in sa_inspect(engine).get_columns("knowledge_documents")
                }
                if "object_key" not in columns:
                    connection.execute(
                        text("ALTER TABLE knowledge_documents ADD COLUMN object_key VARCHAR(1024)")
                    )
                FILE_DATA_NOT_NULL = not columns["file_data"]["nullable"]
        DATABASE_READY = True
        DATABASE_INIT_ERROR = ""
        logger.info("Database initialization completed successfully.")
        return True
    except Exception as exc:
        DATABASE_READY = False
        DATABASE_INIT_ERROR = f"{type(exc).__name__}: {exc}"
        logger.exception("Database initialization failed.")
        return False

async def initialize_database_background() -> None:
    """Retry database initialization in the background so the HTTP port can open."""
    while not DATABASE_READY:
        if await asyncio.to_thread(initialize_database):
            return
        logger.warning("Database is not ready; retrying initialization in 10 seconds.")
        await asyncio.sleep(10)

serializer = URLSafeTimedSerializer(SESSION_SECRET, salt="robot-competition-kb-session") if SESSION_SECRET else None

app = FastAPI(
    title="机器人竞赛知识库",
    version=APP_VERSION,
    description="仅限授权用户上传，上传后自动入库，并为 V4.1.0 智能体提供检索 API。",
)


@app.on_event("startup")
async def start_database_initialization() -> None:
    # Do not wait for PostgreSQL migrations before binding Render's HTTP port.
    asyncio.create_task(initialize_database_background())


ALLOWED_EXTENSIONS = {
    ".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".txt", ".md", ".json",
    ".png", ".jpg", ".jpeg", ".webp"
}


def utcnow() -> datetime:
    return datetime.now(timezone.utc)


def store_raw_file(raw: bytes, digest: str, filename: str, content_type: str) -> tuple[bytes | None, str | None]:
    """Use configured S3-compatible object storage; otherwise retain bytes in the database."""
    if OBJECT_STORAGE_ENABLED and OBJECT_STORAGE_CLIENT:
        safe_name = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(filename).name)[:180] or "file"
        key = f"documents/{digest}/{safe_name}"
        OBJECT_STORAGE_CLIENT.put_object(
            Bucket=OBJECT_STORAGE_BUCKET_NAME, Key=key, Body=raw,
            ContentType=content_type or "application/octet-stream"
        )
        # Old SQLite schemas may still require file_data; duplicate only for that legacy case.
        return (raw if FILE_DATA_NOT_NULL else None), key
    return raw, None


def read_raw_file(doc: "KnowledgeDocument") -> bytes:
    if doc.object_key:
        if not (OBJECT_STORAGE_ENABLED and OBJECT_STORAGE_CLIENT):
            raise RuntimeError("该文件存储在对象存储中，但当前服务尚未配置对象存储凭据。")
        response = OBJECT_STORAGE_CLIENT.get_object(Bucket=OBJECT_STORAGE_BUCKET_NAME, Key=doc.object_key)
        return response["Body"].read()
    if doc.file_data is None:
        raise RuntimeError("文件内容不存在。")
    return bytes(doc.file_data)


def delete_raw_file(object_key: str | None) -> None:
    if object_key and OBJECT_STORAGE_ENABLED and OBJECT_STORAGE_CLIENT:
        OBJECT_STORAGE_CLIENT.delete_object(Bucket=OBJECT_STORAGE_BUCKET_NAME, Key=object_key)


def configured_users() -> dict[str, dict[str, str]]:
    """One shared login for the owner and teacher."""
    username = os.getenv("SHARED_USERNAME", "").strip() or os.getenv("ADMIN_USERNAME", "").strip()
    password = os.getenv("SHARED_PASSWORD", "") or os.getenv("ADMIN_PASSWORD", "")
    if username and password:
        return {username: {"password": password, "role": "admin"}}
    return {}


def current_user(request: Request) -> dict[str, str] | None:
    if not serializer:
        return None
    token = request.cookies.get(COOKIE_NAME, "")
    if not token:
        return None
    try:
        data = serializer.loads(token, max_age=12 * 60 * 60)
        username = str(data.get("username", ""))
        role = str(data.get("role", ""))
        user = configured_users().get(username)
        if user and user["role"] == role:
            return {"username": username, "role": role}
    except (BadSignature, SignatureExpired, ValueError, TypeError):
        return None
    return None


def require_database_ready() -> None:
    if not DATABASE_READY:
        if DATABASE_INIT_ERROR:
            detail = "数据库暂不可用，请稍后重试；管理员请检查 Render 服务日志。"
        else:
            detail = "数据库正在初始化，请稍后重试。"
        raise HTTPException(status_code=503, detail=detail)


def require_user(request: Request) -> dict[str, str]:
    user = current_user(request)
    if not user:
        raise HTTPException(status_code=401, detail="请先登录知识库")
    require_database_ready()
    return user


def authorized_for_search(request: Request) -> bool:
    user = current_user(request)
    if user:
        return True
    if not KB_API_KEY:
        return False
    supplied = request.headers.get("authorization", "")
    if not supplied.lower().startswith("bearer "):
        return False
    token = supplied[7:].strip()
    return hmac.compare_digest(token, KB_API_KEY)


def render_login(message: str = "") -> HTMLResponse:
    msg = f'<div class="msg">{message}</div>' if message else ""
    html = """<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>机器人竞赛知识库登录</title><style>
:root{font-family:Inter,"Microsoft YaHei",Arial,sans-serif;color:#172033;background:#f3f6fb}
*{box-sizing:border-box}body{margin:0;min-height:100vh;display:grid;place-items:center;padding:20px}
.card{width:min(430px,100%);background:#fff;border:1px solid #e2e8f0;border-radius:20px;padding:30px;box-shadow:0 16px 45px #0f172a12}
.mark{width:52px;height:52px;border-radius:15px;background:#e0edff;color:#1d4ed8;display:grid;place-items:center;font-size:25px}
h1{font-size:23px;margin:18px 0 8px}.sub{color:#64748b;font-size:14px;line-height:1.7}
label{display:block;font-size:13px;font-weight:700;margin:18px 0 7px}
input{width:100%;padding:12px;border:1px solid #cbd5e1;border-radius:10px;font-size:15px}
button{margin-top:22px;width:100%;border:0;border-radius:10px;padding:13px;background:#1d4ed8;color:#fff;font-size:15px;font-weight:700;cursor:pointer}
.msg{margin-top:15px;background:#fff1f2;color:#be123c;padding:10px;border-radius:9px;font-size:13px}
</style></head><body><main class="card"><div class="mark">⌘</div><h1>机器人竞赛知识库</h1>
<div class="sub">仅供授权的管理员与老师使用。上传后自动整理入库，无需审核。</div>
""" + msg + """
<form method="post" action="/login">
<label for="username">账号</label><input id="username" name="username" autocomplete="username" required>
<label for="password">密码</label><input id="password" type="password" name="password" autocomplete="current-password" required>
<button type="submit">登录知识库</button></form></main></body></html>"""
    return HTMLResponse(html)


DASHBOARD_HTML = r"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>机器人竞赛知识库</title><style>
:root{font-family:Inter,"Microsoft YaHei",Arial,sans-serif;color:#172033;background:#f4f7fb}
*{box-sizing:border-box}body{margin:0}.top{background:#fff;border-bottom:1px solid #e5eaf2;padding:16px 24px;display:flex;align-items:center;justify-content:space-between;gap:12px;position:sticky;top:0;z-index:5}
.brand{display:flex;align-items:center;gap:12px}.logo{width:42px;height:42px;border-radius:13px;background:#e0edff;color:#1d4ed8;display:grid;place-items:center;font-size:22px}
h1{font-size:19px;margin:0 0 3px}.muted{color:#64748b;font-size:13px}.user{display:flex;align-items:center;gap:12px;font-size:13px}
button,.btn{border:0;border-radius:10px;padding:10px 14px;font-size:13px;font-weight:700;cursor:pointer;background:#1d4ed8;color:#fff}
button.secondary{background:#e8eef8;color:#334155}button.danger{background:#fff1f2;color:#be123c}
main{max-width:1180px;margin:24px auto;padding:0 18px}.grid{display:grid;grid-template-columns:1.2fr .8fr;gap:18px}
.card{background:#fff;border:1px solid #e2e8f0;border-radius:16px;padding:20px;min-width:0}.card h2{font-size:17px;margin:0 0 7px}.help{font-size:13px;color:#64748b;line-height:1.65;margin-bottom:15px}
.drop{border:1.5px dashed #93b4e8;border-radius:14px;background:#f8fbff;padding:26px 16px;text-align:center}
.drop strong{display:block;margin:8px 0}.drop input{max-width:100%;margin-top:12px}
.row{display:flex;gap:9px;align-items:center;flex-wrap:wrap}.stats{display:grid;grid-template-columns:repeat(3,1fr);gap:10px;margin-bottom:18px}
.stat{padding:14px;background:#f8fafc;border-radius:12px}.stat b{display:block;font-size:22px;margin-bottom:3px}.stat span{font-size:12px;color:#64748b}
input[type=text]{padding:11px;border:1px solid #cbd5e1;border-radius:10px;font-size:14px;min-width:0;flex:1}
.status{font-size:13px;color:#475569;margin-top:12px;white-space:pre-wrap}.ok{color:#047857}.err{color:#be123c}
.tablewrap{overflow:auto;margin-top:12px}table{width:100%;border-collapse:collapse;font-size:13px;min-width:600px}th,td{text-align:left;padding:12px 10px;border-bottom:1px solid #edf0f5;vertical-align:top}th{color:#64748b;font-weight:600;background:#fafbfd}
.file{font-weight:700;word-break:break-word}.pill{display:inline-block;padding:4px 8px;background:#dcfce7;color:#166534;border-radius:999px;font-size:11px}
.results{display:grid;gap:10px;margin-top:14px}.result{padding:13px;border:1px solid #e2e8f0;border-radius:12px}.result b{display:block;margin-bottom:5px}.result p{font-size:13px;line-height:1.7;color:#475569;margin:0}.result small{color:#64748b}
.foot{margin:18px 2px;color:#94a3b8;font-size:12px}
@media(max-width:800px){.grid{grid-template-columns:1fr}.top{padding:12px 15px}main{margin:15px auto}.card{padding:16px}.user span{display:none}}
</style></head><body>
<header class="top"><div class="brand"><div class="logo">⌘</div><div><h1>机器人竞赛知识库</h1><div class="muted">上传即入库 · 自动建立检索索引</div></div></div>
<div class="user"><span id="who">已登录</span><form method="post" action="/logout"><button class="secondary" type="submit">退出登录</button></form></div></header>
<main>
<div class="stats"><div class="stat"><b id="docCount">—</b><span>已存储资料</span></div><div class="stat"><b id="chunkCount">—</b><span>检索文本片段</span></div><div class="stat"><b id="statusCount">—</b><span>等待处理 / 无文本资料</span></div></div>
<div class="grid">
<section class="card"><h2>上传新资料</h2><div class="help">支持 PDF、Word、PPT、Excel、CSV、TXT、Markdown 及常见图片。单个文件最大 25 MB，每次最多 30 个文件。上传后直接入库，不需要审核。</div>
<div class="drop"><div style="font-size:30px">⇧</div><strong>选择一个或多个文件</strong><div class="muted">建议上传原始报告、获奖名单、竞赛通知和评分标准</div><input id="files" type="file" multiple accept=".pdf,.docx,.pptx,.xlsx,.csv,.txt,.md,.json,.png,.jpg,.jpeg,.webp"></div>
<div class="row" style="margin-top:13px"><button id="uploadBtn" onclick="uploadFiles()">上传并自动入库</button><button class="secondary" onclick="loadData()">刷新资料列表</button></div><div class="status" id="uploadStatus"></div></section>
<section class="card"><h2>检索知识库</h2><div class="help">输入创意、功能或技术关键词，查看正式知识库中最相关的资料。此检索接口也供 V4.1.0 智能体调用。</div>
<div class="row"><input id="query" type="text" placeholder="例如：儿童接送、UWB、路径规划"><button onclick="searchKb()">检索</button></div><div class="status" id="searchStatus"></div><div class="results" id="results"></div></section>
</div>
<section class="card" style="margin-top:18px"><h2>资料列表</h2><div class="help">重复文件会自动跳过。删除资料后，对应的检索片段也会一并删除。</div><div class="tablewrap"><table><thead><tr><th>文件名称</th><th>上传者</th><th>状态</th><th>上传时间</th><th>操作</th></tr></thead><tbody id="docs"><tr><td colspan="5">正在加载…</td></tr></tbody></table></div></section>
<div class="foot">内部资料库 · 请勿上传未经授权的个人敏感信息或不允许共享的文件。</div>
</main>
<script>
async function api(url,options={}){const r=await fetch(url,options);if(r.status===401){location.href="/login";throw new Error("登录已过期，请重新登录")}let data={};try{data=await r.json()}catch{}if(!r.ok)throw new Error(data.detail||"请求失败");return data}
function setStatus(id,text,kind=""){const el=document.getElementById(id);el.textContent=text;el.className="status "+kind}
async function loadData(){try{const d=await api("/api/v1/kb/documents");document.getElementById("docCount").textContent=d.total_documents;document.getElementById("chunkCount").textContent=d.total_chunks;document.getElementById("statusCount").textContent=d.needs_attention;document.getElementById("who").textContent=d.username+"（"+(d.role==="admin"?"管理员":"老师")+"）";const tbody=document.getElementById("docs");tbody.replaceChildren();if(!d.documents.length){const tr=document.createElement("tr"),td=document.createElement("td");td.colSpan=5;td.textContent="还没有资料，请先上传";tr.append(td);tbody.append(tr)}for(const doc of d.documents){const tr=document.createElement("tr");const td1=document.createElement("td");td1.className="file";td1.textContent=doc.original_name;const td2=document.createElement("td");td2.textContent=doc.uploaded_by;const td3=document.createElement("td");const pill=document.createElement("span");pill.className="pill";pill.textContent=doc.status;td3.append(pill);const td4=document.createElement("td");td4.textContent=doc.created_at;const td5=document.createElement("td");const row=document.createElement("div");row.className="row";const dl=document.createElement("button");dl.className="secondary";dl.textContent="下载";dl.onclick=()=>location.href="/api/v1/kb/documents/"+doc.id+"/download";const del=document.createElement("button");del.className="danger";del.textContent="删除";del.onclick=()=>deleteDoc(doc.id,doc.original_name);row.append(dl,del);td5.append(row);tr.append(td1,td2,td3,td4,td5);tbody.append(tr)}}catch(e){setStatus("uploadStatus",e.message,"err")}}
async function uploadFiles(){const input=document.getElementById("files");if(!input.files.length){setStatus("uploadStatus","请先选择文件","err");return}const fd=new FormData();for(const f of input.files)fd.append("files",f);const b=document.getElementById("uploadBtn");b.disabled=true;b.textContent="上传处理中…";setStatus("uploadStatus","正在上传、提取文字并建立索引，请不要关闭页面");try{const d=await api("/api/v1/kb/upload",{method:"POST",body:fd});setStatus("uploadStatus",d.message+"\n"+d.items.map(x=>x.name+"："+x.status).join("\n"),"ok");input.value="";await loadData()}catch(e){setStatus("uploadStatus",e.message,"err")}finally{b.disabled=false;b.textContent="上传并自动入库"}}
async function searchKb(){const q=document.getElementById("query").value.trim();if(!q){setStatus("searchStatus","请输入检索内容","err");return}setStatus("searchStatus","正在检索…");const box=document.getElementById("results");box.replaceChildren();try{const d=await api("/api/v1/knowledge/search",{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify({query:q,top_k:8})});setStatus("searchStatus","找到 "+d.results.length+" 条相关资料", "ok");for(const r of d.results){const item=document.createElement("div");item.className="result";const title=document.createElement("b");title.textContent=r.title;const p=document.createElement("p");p.textContent=r.content;const sm=document.createElement("small");sm.textContent="来源："+r.source+" · 相关度："+Math.round((r.similarity||0)*100)+"%";item.append(title,p,sm);box.append(item)}}catch(e){setStatus("searchStatus",e.message,"err")}}
async function deleteDoc(id,name){if(!confirm("确定删除“"+name+"”？删除后该资料将不再被智能体检索。"))return;try{await api("/api/v1/kb/documents/"+id,{method:"DELETE"});await loadData()}catch(e){alert(e.message)}}
loadData();
</script></body></html>"""


def extract_text(filename: str, raw: bytes) -> tuple[str, str]:
    ext = Path(filename).suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(raw))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
            return text.strip(), "已入库" if text.strip() else "已保存（需OCR）"
        if ext == ".docx":
            from docx import Document
            doc = Document(io.BytesIO(raw))
            pieces = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    pieces.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(pieces).strip()
            return text, "已入库" if text else "已保存（无可提取文字）"
        if ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(io.BytesIO(raw))
            pieces = []
            for i, slide in enumerate(prs.slides, 1):
                slide_text = [shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
                if slide_text:
                    pieces.append(f"第{i}页：\n" + "\n".join(slide_text))
            text = "\n".join(pieces).strip()
            return text, "已入库" if text else "已保存（无可提取文字）"
        if ext == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            pieces = []
            for ws in wb.worksheets:
                pieces.append(f"工作表：{ws.title}")
                for row in ws.iter_rows(values_only=True):
                    vals = [str(v) for v in row if v is not None and str(v).strip()]
                    if vals:
                        pieces.append(" | ".join(vals))
            text = "\n".join(pieces).strip()
            return text, "已入库" if text else "已保存（无可提取文字）"
        if ext in {".txt", ".md", ".csv", ".json"}:
            text = raw.decode("utf-8-sig", errors="replace").strip()
            return text, "已入库" if text else "已保存（无可提取文字）"
        if ext in {".png", ".jpg", ".jpeg", ".webp"}:
            return "", "已保存（图片OCR待接入）"
    except Exception as exc:
        return "", "解析失败：" + type(exc).__name__
    return "", "已保存（未建立文本索引）"


def chunk_text(text: str, size: int = 900, overlap: int = 120) -> list[str]:
    clean = re.sub(r"[ \t]+", " ", text or "").strip()
    if not clean:
        return []
    chunks = []
    start = 0
    while start < len(clean):
        end = min(len(clean), start + size)
        piece = clean[start:end].strip()
        if piece:
            chunks.append(piece)
        if end >= len(clean):
            break
        start = max(start + 1, end - overlap)
    return chunks


def terms(text: str) -> set[str]:
    s = (text or "").lower()
    out = set(re.findall(r"[a-z0-9_+-]{2,}", s))
    for seq in re.findall(r"[\u4e00-\u9fff]+", s):
        if len(seq) == 1:
            out.add(seq)
        else:
            out.update(seq[i:i+2] for i in range(len(seq)-1))
            if len(seq) <= 8:
                out.add(seq)
    return out


def perform_search(query: str, top_k: int = 5) -> list[dict[str, Any]]:
    require_database_ready()
    query = (query or "").strip()
    if not query:
        return []
    q_terms = terms(query)
    if not q_terms:
        return []
    with SessionLocal() as db:
        rows = db.execute(
            select(KnowledgeChunk, KnowledgeDocument)
            .join(KnowledgeDocument, KnowledgeChunk.document_id == KnowledgeDocument.id)
        ).all()
    best: dict[int, dict[str, Any]] = {}
    for chunk, doc in rows:
        d_terms = terms(chunk.content)
        overlap = q_terms.intersection(d_terms)
        if not overlap:
            continue
        score = len(overlap) / max(1, len(q_terms)) * (0.65 + 0.35 * len(overlap) / max(1, len(d_terms) ** 0.5))
        score = min(1.0, score)
        previous = best.get(doc.id)
        if previous is None or score > previous["similarity"]:
            best[doc.id] = {
                "title": doc.original_name,
                "content": chunk.content[:1400],
                "similarity": round(score, 4),
                "source": doc.original_name,
                "doc_id": doc.id,
                "uploaded_by": doc.uploaded_by,
                "created_at": doc.created_at.isoformat() if doc.created_at else "",
            }
    return sorted(best.values(), key=lambda x: x["similarity"], reverse=True)[:max(1, min(top_k, 20))]


class SearchRequest(BaseModel):
    query: str = Field(..., min_length=1, max_length=4000)
    top_k: int = Field(default=5, ge=1, le=20)


@app.get("/", response_class=HTMLResponse)
def root(request: Request):
    if current_user(request):
        return HTMLResponse(DASHBOARD_HTML)
    return RedirectResponse("/login", status_code=303)


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    if current_user(request):
        return RedirectResponse("/", status_code=303)
    return render_login()


@app.post("/login")
def login(request: Request, username: str = Form(...), password: str = Form(...)):
    users = configured_users()
    if not serializer:
        return render_login("服务尚未配置 SESSION_SECRET，请管理员在部署环境中设置。")
    user = users.get(username.strip())
    if not user or not hmac.compare_digest(user["password"], password):
        return render_login("账号或密码不正确。")
    token = serializer.dumps({"username": username.strip(), "role": user["role"]})
    response = RedirectResponse("/", status_code=303)
    response.set_cookie(
        COOKIE_NAME, token, max_age=12 * 60 * 60, httponly=True,
        secure=COOKIE_SECURE, samesite="strict", path="/"
    )
    return response


@app.post("/logout")
def logout():
    response = RedirectResponse("/login", status_code=303)
    response.delete_cookie(COOKIE_NAME, path="/")
    return response


@app.get("/health")
def health():
    # Keep the health endpoint responsive while the database initializes/retries.
    base = {
        "ok": True,
        "service": "robot-competition-knowledge-base",
        "version": APP_VERSION,
        "database": "postgresql" if DATABASE_URL else "sqlite",
        "database_ready": DATABASE_READY,
        "object_storage": OBJECT_STORAGE_PROVIDER if OBJECT_STORAGE_ENABLED else "database_blob",
    }
    if not DATABASE_READY:
        base["database_status"] = "unavailable" if DATABASE_INIT_ERROR else "initializing"
        return base
    with SessionLocal() as db:
        count = db.scalar(select(func.count()).select_from(KnowledgeDocument)) or 0
        chunks = db.scalar(select(func.count()).select_from(KnowledgeChunk)) or 0
    base.update({"documents": count, "chunks": chunks, "database_status": "ready"})
    return base


@app.get("/api/v1/kb/documents")
def list_documents(request: Request):
    user = require_user(request)
    with SessionLocal() as db:
        docs = db.scalars(select(KnowledgeDocument).order_by(KnowledgeDocument.created_at.desc())).all()
        chunk_count = db.scalar(select(func.count()).select_from(KnowledgeChunk)) or 0
        attention = sum(1 for d in docs if not d.extracted_text.strip())
        items = [{
            "id": d.id, "original_name": d.original_name, "status": d.status,
            "uploaded_by": d.uploaded_by,
            "created_at": d.created_at.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M UTC") if d.created_at else "",
            "has_text": bool(d.extracted_text.strip()),
        } for d in docs]
    return {"username": user["username"], "role": user["role"], "total_documents": len(items),
            "total_chunks": chunk_count, "needs_attention": attention, "documents": items}


@app.post("/api/v1/kb/upload")
async def upload_documents(request: Request, files: list[UploadFile] = File(...)):
    user = require_user(request)
    if not files:
        raise HTTPException(status_code=400, detail="请先选择文件")
    if len(files) > MAX_UPLOADS:
        raise HTTPException(status_code=413, detail=f"一次最多上传 {MAX_UPLOADS} 个文件")
    results = []
    for upload in files:
        filename = Path(upload.filename or "unnamed").name
        ext = Path(filename).suffix.lower()
        if ext not in ALLOWED_EXTENSIONS:
            results.append({"name": filename, "status": "不支持的文件类型"})
            continue
        raw = await upload.read()
        if len(raw) > MAX_FILE_BYTES:
            results.append({"name": filename, "status": "文件超过 25 MB 限制"})
            continue
        if not raw:
            results.append({"name": filename, "status": "空文件，已跳过"})
            continue
        digest = hashlib.sha256(raw).hexdigest()
        with SessionLocal() as db:
            existing = db.scalar(select(KnowledgeDocument).where(KnowledgeDocument.sha256 == digest))
            if existing:
                results.append({"name": filename, "status": f"重复文件，已跳过（已有：{existing.original_name}）"})
                continue
        extracted, status = extract_text(filename, raw)
        try:
            file_data, object_key = store_raw_file(raw, digest, filename, upload.content_type or "application/octet-stream")
        except Exception as exc:
            results.append({"name": filename, "status": "文件存储失败：" + type(exc).__name__})
            continue
        doc = KnowledgeDocument(
            original_name=filename[:255], sha256=digest,
            content_type=upload.content_type or "application/octet-stream",
            file_data=file_data, object_key=object_key, extracted_text=extracted, status=status,
            uploaded_by=user["username"], created_at=utcnow(),
        )
        doc.chunks = [KnowledgeChunk(chunk_index=i, content=piece) for i, piece in enumerate(chunk_text(extracted))]
        with SessionLocal() as db:
            db.add(doc)
            try:
                db.commit()
                results.append({"name": filename, "status": status, "document_id": doc.id, "indexed_chunks": len(doc.chunks)})
            except Exception:
                db.rollback()
                delete_raw_file(object_key)
                results.append({"name": filename, "status": "保存失败，请检查数据库连接或日志"})
    ok = sum(1 for x in results if "document_id" in x)
    return {"success": True, "uploaded": ok, "items": results,
            "message": f"处理完成：成功入库 {ok} 个文件；未入库或跳过 {len(results)-ok} 个。"}


@app.delete("/api/v1/kb/documents/{document_id}")
def delete_document(document_id: int, request: Request):
    require_user(request)
    with SessionLocal() as db:
        doc = db.get(KnowledgeDocument, document_id)
        if not doc:
            raise HTTPException(status_code=404, detail="资料不存在")
        object_key = doc.object_key
        db.delete(doc)
        db.commit()
    delete_raw_file(object_key)
    return {"success": True, "deleted_id": document_id}


@app.get("/api/v1/kb/documents/{document_id}/download")
def download_document(document_id: int, request: Request):
    require_user(request)
    with SessionLocal() as db:
        doc = db.get(KnowledgeDocument, document_id)
        if not doc:
            raise HTTPException(status_code=404, detail="资料不存在")
        try:
            raw = read_raw_file(doc)
        except Exception as exc:
            raise HTTPException(status_code=503, detail="无法读取原始文件，请检查对象存储配置。") from exc
        name = doc.original_name.replace('"', "")
        mime = doc.content_type
    return Response(content=raw, media_type=mime, headers={
        "Content-Disposition": "attachment; filename*=UTF-8''" + __import__("urllib.parse").parse.quote(name)
    })


@app.post("/api/v1/knowledge/search")
def knowledge_search(payload: SearchRequest, request: Request):
    if not authorized_for_search(request):
        raise HTTPException(status_code=401, detail="未授权。请登录管理页面或配置正确的 KB_API_KEY。")
    results = perform_search(payload.query, payload.top_k)
    sources = list(dict.fromkeys(x["source"] for x in results))
    return {"success": True, "query": payload.query, "results": results, "cases": results, "sources": sources}


@app.get("/api/v1/knowledge/search")
def knowledge_search_get(request: Request, q: str, top_k: int = 5):
    if not authorized_for_search(request):
        raise HTTPException(status_code=401, detail="未授权。")
    results = perform_search(q, top_k)
    return {"success": True, "query": q, "results": results, "cases": results,
            "sources": list(dict.fromkeys(x["source"] for x in results))}
